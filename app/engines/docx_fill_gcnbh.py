"""
Thay chuoi trong file .docx ma KHONG lam vo dinh dang.

Van de: Word cat mot cau thanh nhieu "run" theo lich su go/sua chinh ta.
Vi du trong giay chung nhan mau:
    "So: 26/45/07/TSTH/C000021/15/01"  ->  ['26/45/07/TSTH/C000021/15', '/0', '1']
    "Tu 02/07/2026 den 02/07/2027"     ->  10 run roi
Nen tim-va-thay tren tung run se truot het. Con gop het run cua doan lai lam
mot thi mat dinh dang rieng cua tung phan (vd dong tieu de co 3 co chu khac nhau).

Cach o day: xac dinh DAI RUN phu len doan can thay, viet ket qua vao run DAU
cua dai do va xoa phan bi phu o cac run sau. Chu moi thay se mang dinh dang
cua run dau - dung y muon, vi cac run bi cat vun deu cung dinh dang.
"""

from __future__ import annotations

import docx


def replace_in_paragraph(p, old: str, new: str) -> bool:
    """Thay lan xuat hien dau tien cua `old` trong doan `p`. Tra ve True neu co thay."""
    runs = p.runs
    if not runs:
        return False
    full = "".join(r.text for r in runs)
    at = full.find(old)
    if at < 0:
        return False
    end = at + len(old)

    # ban do vi tri: run i phu doan [start_i, stop_i)
    spans, pos = [], 0
    for r in runs:
        spans.append((pos, pos + len(r.text)))
        pos += len(r.text)

    first = None
    for i, (s, e) in enumerate(spans):
        if e <= at or s >= end:
            continue                      # run nam ngoai vung thay
        head = runs[i].text[: max(0, at - s)]
        tail = runs[i].text[max(0, end - s):]
        if first is None:
            first = i
            runs[i].text = head + new + tail
        else:
            runs[i].text = tail           # phan bi phu da nam o run dau
    return first is not None


def iter_paragraphs(doc):
    """Duyet moi doan: than bai, trong bang (ke ca bang long nhau), header/footer."""
    yield from doc.paragraphs
    for t in doc.tables:
        yield from _iter_table(t)
    for sec in doc.sections:
        for part in (sec.header, sec.footer):
            yield from part.paragraphs
            for t in part.tables:
                yield from _iter_table(t)


def _iter_table(t):
    for row in t.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for inner in cell.tables:
                yield from _iter_table(inner)


def replace_all(doc, mapping: dict) -> dict:
    """Thay moi cap trong `mapping` tren toan bo tai lieu. Tra ve so lan thay duoc."""
    hits = {k: 0 for k in mapping}
    for p in iter_paragraphs(doc):
        for old, new in mapping.items():
            # mot doan co the chua cung mot khoa nhieu lan
            while replace_in_paragraph(p, old, str(new)):
                hits[old] += 1
    return hits


def load(path):
    return docx.Document(path)


# ---------------------------------------------------------------------------
# Bo sung cho GCNBH TBHH: thay theo bieu thuc chinh quy, va don comment cua Word.

import re as _re

from docx.oxml.ns import qn as _qn


def replace_regex(doc, pattern: str, repl, only_if: str = "") -> int:
    """Thay lan xuat hien dau tien cua `pattern` trong moi doan.

    Dung cho dong "Ho Chi Minh, ngay .. thang .. nam .." - cho trong nay khong
    co dau ngoac [..] nen khong thay theo chuoi co dinh duoc. `only_if` la chuoi
    bat buoc phai co trong doan (chan an nham cac cau khac cung chua chu "ngay").
    """
    rx = _re.compile(pattern)
    n = 0
    for p in iter_paragraphs(doc):
        text = p.text
        if only_if and only_if not in text:
            continue
        m = rx.search(text)
        if m and replace_in_paragraph(p, m.group(0), repl):
            n += 1
    return n


def strip_comments(doc) -> None:
    """Bo chu thich Word khoi ban xuat (mau co san mot comment '[Ngay cap don]')."""
    root = doc.element
    for tag in ("w:commentRangeStart", "w:commentRangeEnd"):
        for el in root.findall(".//" + _qn(tag)):
            el.getparent().remove(el)
    for ref in root.findall(".//" + _qn("w:commentReference")):
        run = ref.getparent()
        run.getparent().remove(run)


def strip_highlight(doc) -> int:
    """Bo boi vang khoi ban xuat.

    Trong ca hai mau, moi cho boi vang deu la cho danh dau [..] (20 cho o mau
    goc, 25 cho o mau SDBS) - khong co doan noi dung nao boi vang that, va
    styles.xml cung khong dinh nghia highlight. Nen quet sach ca tai lieu la an
    toan; chu danh may vao van giu nguyen font/co chu cua run cu.
    """
    n = 0
    parts = [doc.element]
    for sec in doc.sections:
        parts += [sec.header._element, sec.footer._element]
    for root in parts:
        for hl in root.findall(".//" + _qn("w:highlight")):
            hl.getparent().remove(hl)
            n += 1
    return n
