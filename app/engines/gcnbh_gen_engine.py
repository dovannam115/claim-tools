# -*- coding: utf-8 -*-
"""Sinh GIAY CHUNG NHAN BAO HIEM (ban goc) va SUA DOI BO SUNG (SDBS) tu mot file
Excel xuat tu PIAS - moi dong mot giay - bang cach dien vao cac file mau Word.

Ruot lay tu tool .bat Documents\\GCNBH-TBHH (ban do ngung phat cho user tu
22/08/2026, web nay la ban duy nhat). Khac ban .bat:
  - bo hop thoai chon file + thanh tien trinh console, thay bang callback cho
    Streamlit;
  - Word (COM) chay trong luong phu cua Streamlit nen phai CoInitialize, va
    khoa lai vi web co the co hai nguoi bam cung luc;
  - them truong [Ngay cap don goc] tren mau SDBS (xem resolve_goc_dates).

Mau Word de o gcnbh_templates\\ canh file nay:
    GCNBH_<MA SP>_Template.docx   ->  ban goc cua san pham do
    file nao ten co chu SDBS      ->  mau sua doi bo sung (dung chung moi sp)
Them san pham moi = tha them mot mau vao thu muc do, KHONG phai sua code. Mau
moi lay tu mot giay that thi chay Documents\\GCNBH-TBHH\\tool\\make_template.py
de doi cac cho boi vang thanh cho danh dau [..].

Ghi chu ve du lieu:
  - Do cot theo TEN tieu de (khong theo vi tri) vi thu tu cot hai sheet lech nhau.
  - Cot "Phi bao hiem - VAT" la phi TRUOC thue, cot "Phi BH" la phi DA GOM thue
    (kiem chung: 27.969.364 + 2.796.936 = 30.766.300). Ca mau ban goc lan mau
    SDBS deu lay cot "Phi BH" (Nam chot 28/10/2026: cho [Phi BH] cua SDBS la cot
    Phi BH, khong phai cot - VAT).
  - Sheet SDBS de so am cho ky giam. Mau da co chu tang/giam o cho [SDBS] nen
    so tien in ra lay tri tuyet doi.
  - Mau khong co cho danh dau so thang chuong trinh -> suy ra tu khoang cach
    Ngay bat dau -> Ngay ket thuc. Muon chot tay thi them cot "Thang BH".
"""

from __future__ import annotations

import copy
import datetime as dt
import glob
import io
import os
import re
import threading

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook

import docx_fill_gcnbh as docx_fill

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATES = os.path.join(HERE, "gcnbh_templates")

# dau phan cach hang nghin - cac giay that cua phong dung dau cham
THOUSANDS = "."


class ToolError(Exception):
    pass


# --------------------------------------------------------------- tien ich nho

def norm(v) -> str:
    s = str(v if v is not None else "").replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip().lower()


def fold(v) -> str:
    """Chuan hoa de so ten: bo dau D gach ngang, ve chu thuong."""
    return norm(v).replace("đ", "d")


def txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def money(v) -> str:
    try:
        s = format(round(abs(float(v))), ",")
    except (TypeError, ValueError):
        return txt(v)
    return s if THOUSANDS == "," else s.replace(",", THOUSANDS)


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    m = re.match(r"\s*(\d{1,2})[/-](\d{1,2})[/-](\d{4})", txt(v))
    if m:
        return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


def fdate(v) -> str:
    d = as_date(v)
    return d.strftime("%d/%m/%Y") if d else txt(v)


def months_between(tu, den):
    """So thang chuong trinh + co bao 'lam tron khong chac' (True = nen kiem tay)."""
    d1, d2 = as_date(tu), as_date(den)
    if not d1 or not d2:
        return None, True
    raw = (d2 - d1).days / 30.4375
    n = max(1, round(raw))
    return n, abs(raw - n) > 0.30


def safe_name(s) -> str:
    s = re.sub(r"[\\/:*?<>|\"]+", "_", txt(s)).strip(" ._")
    return s or "no-number"


# --------------------------------------------------------- tim va doc file mau
#
# Ten san pham lay tu TEN FILE MAU: GCNBH_BHMR_Template.docx -> BHMR.
# File mau nao ten co chu SDBS thi la mau sua doi bo sung (dung chung moi sp).

NOISE = {"gcnbh", "gcn", "template", "mau", "mẫu", "giay", "giấy", "chung", "chứng",
         "nhan", "nhận", "bao", "bảo", "hiem", "hiểm", "bh", "copy", "final",
         "new", "moi", "mới", "ban", "bản"}


def _noise_tok(t) -> bool:
    """Chu 'template' hay bi go sai (da gap 'Templatre') nen bat ca ho templ*,
    khong thi ma san pham se dinh them chu do."""
    n = norm(t)
    return n in NOISE or n.startswith("templ")


def product_code(tpl_path) -> str:
    base = os.path.splitext(os.path.basename(tpl_path))[0]
    toks = [t for t in re.split(r"[\s_\-.]+", base) if t]
    keep = [t for t in toks if not _noise_tok(t)]
    return " ".join(keep)


def find_templates(folder=TEMPLATES):
    """Tra ve ({ma sp: duong dan mau goc}, duong dan mau SDBS)."""
    gcn, sdbs = {}, ""
    for p in sorted(glob.glob(os.path.join(folder, "*.docx"))):
        name = os.path.basename(p)
        if name.startswith("~$"):
            continue
        if "sdbs" in fold(name):
            sdbs = sdbs or p
            continue
        code = product_code(p)
        if code:
            gcn.setdefault(code, p)
    if not gcn:
        raise ToolError(
            "Khong tim thay mau GCNBH nao trong %s. Can file dat ten kieu "
            "GCNBH_<MA SP>_Template.docx." % folder)
    return gcn, sdbs


# ------------------------------------------------------- dat ten file xuat ra
#
# Luat: <ma san pham> <6 ky tu cuoi so don>.
#   26/45/09/KHAC/P000925      ->  "BHMR 000925"
#   26/45/09/KHAC/P000696/E01  ->  "BHMR 000696 E01"   (6 so cuoi cua DON GOC)

E_TAIL = re.compile(r"[/_\-\s]*(E\s*\d{1,3})\s*$", re.I)


def tail6(so_don) -> str:
    """6 ky tu cuoi cua so don, bo dau phan cach neu lo dinh vao dau."""
    s = txt(so_don)
    t = re.sub(r"^[^0-9A-Za-zÀ-ỹ]+", "", s[-6:])
    return t or s


def name_goc(prod, rec) -> str:
    return "%s %s" % (prod, tail6(rec["so_don"]))


def name_sdbs(prod, rec) -> str:
    so = txt(rec["so_don"])
    m = E_TAIL.search(so)
    goc = txt(rec["so_goc"]) or (so[:m.start()] if m else so)
    out = "%s %s" % (prod, tail6(goc))
    if m:
        out += " " + m.group(1).upper().replace(" ", "")
    return out


# ------------------------------------------------------------- doc file Excel
#
# "ngay_cap" phai loai tru chu "goc": neu file PIAS co ca hai cot "Ngay cap don"
# va "Ngay cap don goc" thi vi tri cot nao dung truoc se an cot kia.

FIELDS = {
    "so_don":       lambda h: h == "số đơn bh/sđbs",
    "so_goc":       lambda h: "số đơn" in h and "gốc" in h,
    "ngay_cap":     lambda h: "ngày cấp" in h and "gốc" not in h,
    "ngay_cap_goc": lambda h: "ngày cấp" in h and "gốc" in h,
    "tu":           lambda h: "ngày bắt đầu" in h,
    "den":          lambda h: "ngày kết thúc" in h,
    "stbh":         lambda h: "số tiền" in h,
    "phi_net":      lambda h: "phí" in h and "vat" in h,
    "phi_gross":    lambda h: h in ("phí bh", "phí bảo hiểm"),
    "so_luong":     lambda h: "số lượng" in h,
    "loai":         lambda h: h == "sđbs",
    "thang":        lambda h: "tháng bh" in h,
}


def sheet_kind(name):
    """Phan loai sheet theo TEN: chi can co chu 'SDBS' hoac 'GCN' la duoc.
    Xet SDBS truoc vi ten kieu 'GCN SDBS' phai tinh la sheet sua doi bo sung."""
    n = fold(name)
    if "sdbs" in n:
        return "sdbs"
    if "gcn" in n:
        return "goc"
    return ""


def tight(v) -> str:
    """Bo het ky tu khong phai chu/so de so ten. Can vi ma san pham nhieu chu
    ('BHNC - ICT') thi ten mau va ten sheet hay ghi dau noi khac nhau:
    mau 'GCNBH_BHNC - ICT_Template.docx' -> ma 'BHNC ICT', con sheet co the la
    'GCN goc BHNC - ICT' hoac 'GCN goc BHNC-ICT'."""
    return re.sub(r"[^0-9a-z]+", "", fold(v))


def sheet_product(name, codes):
    """Ma san pham nam trong ten sheet: 'GCN goc BHMR' -> 'BHMR'.
    Xet ma DAI truoc de 'BHNC ICT' khong bi mot ma ngan hon an mat."""
    n = tight(name)
    for code in sorted(codes, key=lambda c: len(tight(c)), reverse=True):
        if tight(code) in n:
            return code
    return ""


def read_sheet(ws):
    """Tra ve (co cot 'So don BH goc' khong, list[dict] cac dong)."""
    hdr = 0
    for r in range(1, min(ws.max_row, 20) + 1):
        for c in range(1, ws.max_column + 1):
            if norm(ws.cell(r, c).value).startswith("số đơn"):
                hdr = r
                break
        if hdr:
            break
    if not hdr:
        return None, []

    heads = {c: norm(ws.cell(hdr, c).value) for c in range(1, ws.max_column + 1)}
    col = {}
    for key, pred in FIELDS.items():
        col[key] = next((c for c, h in heads.items() if h and pred(h)), 0)
    if not col["so_don"]:
        col["so_don"] = next(
            (c for c, h in heads.items() if h and "số đơn" in h and "gốc" not in h), 0)

    rows = []
    for r in range(hdr + 1, ws.max_row + 1):
        rec = {k: (ws.cell(r, c).value if c else None) for k, c in col.items()}
        if not txt(rec["so_don"]) and not txt(rec["stbh"]):
            continue
        rec["_row"] = r
        rows.append(rec)
    return bool(col["so_goc"]), rows


def check_blank_columns(name, rows, warns):
    """Bay quen thuoc cua openpyxl: doc bang data_only nhung file chua tung duoc
    Excel tinh lai thi o cong thuc tra ve None. Mot cot tien ma TRONG SACH ca
    cot thi gan nhu chac la dinh loi nay chu khong phai du lieu that."""
    for key, label in (("stbh", "Số tiền BH"), ("phi_gross", "Phí BH"),
                       ("phi_net", "Phí bảo hiểm - VAT")):
        vals = [r[key] for r in rows]
        if vals and all(v is None for v in vals):
            warns.append("Sheet '%s': cột '%s' rỗng ở mọi dòng. Nếu cột đó là công "
                         "thức thì mở file bằng Excel, lưu lại một lần rồi tải lên "
                         "lại (openpyxl không tự tính công thức)." % (name, label))


# ------------------------------------------- ngay cap cua DON GOC tren mau SDBS
#
# Cau tren mau SDBS la "Giay chung nhan bao hiem so [..] cap ngay [..]" nen ngay
# o do phai la ngay cap DON GOC. Ban .bat truoc 22/08/2026 dien ngay cap cua
# chinh cai SDBS vao cho nay (chot 18/08 vi luc do khong co nguon nao chac).
# Nay lay theo thu tu: cot "Ngay cap don goc" cua sheet SDBS -> do sang sheet
# GCN goc theo so don goc -> khong ra thi de trong + canh bao (dot test 18/08 co
# 4 don P000534/578/619/620 khong nam trong file nen se roi vao nhanh nay).

def _key(so) -> str:
    return txt(so).upper()


def resolve_goc_dates(sheets, warns):
    """Dien rec['ngay_cap_goc'] cho moi dong SDBS. Tra ve so dong khong tim ra."""
    lookup = {}
    for _name, kind, _prod, rows in sheets:
        if kind != "goc":
            continue
        for r in rows:
            k = _key(r["so_don"])
            if k and r.get("ngay_cap") is not None:
                lookup.setdefault(k, r["ngay_cap"])

    missing = 0
    for name, kind, _prod, rows in sheets:
        if kind != "sdbs":
            continue
        for r in rows:
            if r.get("ngay_cap_goc") not in (None, ""):
                r["_goc_src"] = "cột Excel"
                continue
            hit = lookup.get(_key(r["so_goc"]))
            if hit is not None:
                r["ngay_cap_goc"] = hit
                r["_goc_src"] = "dò sheet GCN gốc"
            else:
                r["_goc_src"] = ""
                missing += 1
                warns.append(
                    "Sheet '%s' dòng %s (%s): không có ngày cấp đơn gốc, đơn gốc %s "
                    "không nằm trong file và cũng không có cột 'Ngày cấp đơn gốc'. "
                    "Giấy in ra để trống chỗ đó."
                    % (name, r["_row"], txt(r["so_don"]), txt(r["so_goc"]) or "?"))
    return missing


def load_input(path, codes, warns):
    """Tra ve list[(ten sheet, loai, ma sp, cac dong)]."""
    wb = load_workbook(path, data_only=True)
    out = []
    for name in wb.sheetnames:
        has_goc_col, rows = read_sheet(wb[name])
        if not rows:
            continue
        check_blank_columns(name, rows, warns)
        kind = sheet_kind(name)
        if not kind:
            kind = "sdbs" if has_goc_col else "goc"
            warns.append("Sheet '%s': tên không có chữ GCN lẫn SĐBS, tool coi là %s."
                         % (name, kind.upper()))
        prod = sheet_product(name, codes)
        if not prod:
            if len(codes) == 1:
                prod = list(codes)[0]
                warns.append("Sheet '%s': tên sheet không có mã sản phẩm, dùng mẫu "
                             "duy nhất đang có (%s)." % (name, prod))
            else:
                raise ToolError(
                    "Sheet '%s': không biết là sản phẩm nào. Ghi mã sản phẩm vào tên "
                    "sheet. Các mẫu đang có: %s." % (name, ", ".join(sorted(codes))))
        out.append((name, kind, prod, rows))
    wb.close()
    if not out:
        raise ToolError("Không thấy dòng dữ liệu nào: không sheet nào có cột "
                        "'Số đơn BH/SĐBS'.")
    resolve_goc_dates(out, warns)
    return out


# ------------------------------------------------------------------ dien mau

DATE_LINE = r"ngày\s*\d*\s*tháng\s*\d*\s*năm\s*\d*"
MUST_HAVE = "[Số đơn BH/SĐBS]"

# Cho danh dau co the vang o mot so mau ma khong phai loi (mau SDBS cu dung
# [Ngay cap don], mau moi dung [Ngay cap don goc]) - vang thi im lang.
OPTIONAL = ("[Ngày cấp đơn gốc]", "[Ngày cấp đơn]")


def _finish(doc, ngay_cap, warns, tag):
    d = as_date(ngay_cap)
    if d:
        stamp = "ngày %02d tháng %02d năm %d" % (d.day, d.month, d.year)
        if not docx_fill.replace_regex(doc, DATE_LINE, stamp, only_if="Chí Minh"):
            warns.append("%s: không tìm thấy dòng 'Hồ Chí Minh, ngày … tháng … "
                         "năm …' trong mẫu." % tag)
    else:
        warns.append("%s: không có ngày cấp đơn, dòng ngày tháng để trống." % tag)
    docx_fill.strip_comments(doc)
    docx_fill.strip_highlight(doc)   # mau boi vang cho danh dau - ban xuat khong boi
    return doc


def _check(hits, tag, tpl, warns, optional=()):
    """Thieu cho danh dau so don thi dung han; thieu cho khac thi chi canh bao,
    vi moi san pham mot mau, co san pham khong co truong do that."""
    missed = [k for k, v in hits.items() if v == 0 and k not in optional]
    if MUST_HAVE in missed:
        raise ToolError(
            "Mẫu %s không có chỗ đánh dấu %s. Mở mẫu ra kiểm tra chính tả trong "
            "ngoặc vuông." % (os.path.basename(tpl), MUST_HAVE))
    if missed:
        warns.append("%s: mẫu %s không có chỗ đánh dấu %s, trường đó không được in."
                     % (tag, os.path.basename(tpl), ", ".join(missed)))


def thang_bh(rec, tag, warns):
    v = rec.get("thang")
    if v not in (None, ""):
        return txt(v)
    n, unsure = months_between(rec["tu"], rec["den"])
    if n is None:
        warns.append("%s: thiếu ngày bắt đầu/kết thúc, không suy ra được số tháng." % tag)
        return ""
    if unsure:
        warns.append("%s: thời hạn lẻ, tool ghi %d tháng, kiểm tra lại." % (tag, n))
    return str(n)


def fill_goc(rec, tpl, warns):
    tag = txt(rec["so_don"]) or "dòng %s" % rec["_row"]
    thang = thang_bh(rec, tag, warns)
    # ban goc phai in phi DA GOM VAT (cot "Phi BH"). Thieu cot do ma lang le lay
    # cot truoc VAT thi giay in ra sai tien duoi dong chu "da bao gom 10% VAT",
    # nen phai keu to.
    phi = rec["phi_gross"]
    if phi is None:
        phi = rec["phi_net"]
        warns.append("%s: cột 'Phí BH' (đã gồm VAT) rỗng, tool in số TRƯỚC VAT, "
                     "KIỂM TRA LẠI." % tag)
    doc = docx_fill.load(tpl)
    hits = docx_fill.replace_all(doc, {
        "[Số đơn BH/SĐBS]": txt(rec["so_don"]),
        "[Tháng BH]":       thang,
        "[Ngày bắt đầu]":   fdate(rec["tu"]),
        "[Ngày kết thúc]":  fdate(rec["den"]),
        "[Số tiền BH]":     money(rec["stbh"]),
        "[Phí BH]":         money(phi),
    })
    _check(hits, tag, tpl, warns)
    return _finish(doc, rec["ngay_cap"], warns, tag)


def fill_sdbs(rec, tpl, warns):
    tag = txt(rec["so_don"]) or "dòng %s" % rec["_row"]
    loai = txt(rec["loai"]).lower()
    if loai not in ("tăng", "giảm"):
        warns.append("%s: cột SĐBS ghi '%s', đáng lẽ phải là tăng hoặc giảm."
                     % (tag, txt(rec["loai"])))
    if not txt(rec["so_luong"]):
        warns.append("%s: thiếu số lượng." % tag)
    phi = rec["phi_gross"]
    if phi is None:
        phi = rec["phi_net"]
        warns.append("%s: cột 'Phí BH' rỗng, tool in số ở cột 'Phí bảo hiểm - VAT', "
                     "KIỂM TRA LẠI." % tag)
    doc = docx_fill.load(tpl)
    hits = docx_fill.replace_all(doc, {
        "[Số đơn BH/SĐBS]":   txt(rec["so_don"]),
        "[Số đơn BH gốc]":    txt(rec["so_goc"]),
        "[Ngày cấp đơn gốc]": fdate(rec.get("ngay_cap_goc")),
        "[Ngày cấp đơn]":     fdate(rec["ngay_cap"]),
        "[Số lượng]":         txt(rec["so_luong"]),
        "[SĐBS]":             loai or txt(rec["loai"]),
        "[Số tiền BH]":       money(rec["stbh"]),
        "[Phí BH]":           money(phi),
    })
    _check(hits, tag, tpl, warns, optional=OPTIONAL)
    # dong ngay thang cuoi trang van la ngay cap cua chinh cai SDBS
    return _finish(doc, rec["ngay_cap"], warns, tag)


# ------------------------------------------------------------------- gop file
#
# Cac mau khong co anh, danh sach dau dong la bullet chu khong phai so thu tu,
# nen gop chi la chep cac phan tu cua <w:body> sang. Nhung co MOT BAY LON:
#
#   Cac mau (nhat la mau SDBS - 18/18 doan) duoc dung bang FRAME NEO VAO TRANG:
#   <w:framePr vAnchor="page" x=.. y=..>. Kieu nay hay gap o file convert tu
#   PDF. Neu chi chen ngat trang (<w:br w:type="page">) thi 15 ban sao van cung
#   mot toa do tuyet doi -> Word don het ve MOT trang, chong len nhau (da thu:
#   270 doan, 2.955 tu, xuat PDF ra dung 1 trang).
#
#   Cach dung: ket thuc moi ban bang mot SECTION BREAK sang trang moi (sectPr
#   dat trong <w:pPr> cua doan cuoi). Moi ban thanh mot section rieng nen frame
#   neo dung vao trang cua no.
#
# Hai thu phu phai xu ly: sectPr cuoi cung luon nam cuoi body (chen truoc no)
# va bookmark se trung ten khi lap lai.

def _strip_bookmarks(el):
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for b in el.findall(".//" + qn(tag)):
            b.getparent().remove(b)


def _section_break(sect_src):
    """Doan rong mang sectPr = ket thuc section hien tai, sang trang moi."""
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    sect = copy.deepcopy(sect_src)
    for t in sect.findall(qn("w:type")):
        sect.remove(t)          # bo 'continuous' -> mac dinh la nextPage
    pPr.append(sect)
    p.append(pPr)
    return p


def merge(docs):
    """Moi giay mot trang, noi tiep nhau trong MOT file .docx."""
    merged = docs[0]
    body = merged.element.body
    tail = body.find(qn("w:sectPr"))
    if tail is None:
        tail = OxmlElement("w:sectPr")
        body.append(tail)
    # sectPr cuoi quyet dinh section CUOI bat dau o dau: mau SDBS de
    # 'continuous' nen ban cuoi cung deo vao trang cua ban ke truoc.
    for t in tail.findall(qn("w:type")):
        tail.remove(t)
    for doc in docs[1:]:
        tail.addprevious(_section_break(tail))
        src_sect = doc.element.body.find(qn("w:sectPr"))
        for el in doc.element.body.iterchildren():
            if el is src_sect:
                continue
            new = copy.deepcopy(el)
            _strip_bookmarks(new)
            tail.addprevious(new)
    return merged


# --------------------------------------------------- ban gop: xuat ra PDF
#
# Ban GOP de in mot luot thi xuat PDF, khong giu .docx. Van phai dung .docx lam
# ban trung gian roi nho Word chuyen (python khong tu dung duoc PDF giu nguyen
# frame/phan trang cua mau) - xong thi xoa .docx.
#
# Khac ban .bat: Streamlit chay script trong LUONG PHU nen phai CoInitialize
# truoc khi goi COM, va phai khoa lai - hai nguoi bam cung luc se mo hai ban
# Word, may de treo. May chu khong co Word/pywin32 thi giu .docx + canh bao chu
# khong lam hong ca lan chay.

_WORD_LOCK = threading.Lock()


def word_available() -> bool:
    """Co du do nghe de xuat PDF khong (de trang web bat/tat o tick cho dung)."""
    try:
        import pythoncom  # noqa: F401
        import win32com.client  # noqa: F401
    except ImportError:
        return False
    return True


def to_pdf(paths, warns, progress=None):
    """Doi cac file .docx gop sang PDF. Tra ve danh sach duong dan cuoi cung."""
    if not paths:
        return []
    try:
        import pythoncom
        import win32com.client as w32
    except ImportError:
        warns.append("Máy chủ không có pywin32 nên không xuất được PDF, giữ .docx.")
        return list(paths)

    made, word = [], None
    with _WORD_LOCK:
        pythoncom.CoInitialize()
        try:
            word = w32.DispatchEx("Word.Application")   # ban Word rieng, khong dung vao file dang mo
            word.Visible = False
            word.DisplayAlerts = 0
            for i, p in enumerate(paths, 1):
                pdf = os.path.splitext(p)[0] + ".pdf"
                if progress:
                    progress(i, len(paths), "PDF: " + os.path.basename(pdf))
                doc = word.Documents.Open(os.path.abspath(p), ReadOnly=True,
                                          AddToRecentFiles=False, ConfirmConversions=False)
                try:
                    doc.ExportAsFixedFormat(os.path.abspath(pdf), 17)   # 17 = wdExportFormatPDF
                finally:
                    doc.Close(False)
                made.append((p, pdf))
        except Exception as exc:
            warns.append("Xuất PDF hỏng (%s), giữ bản .docx." % exc)
        finally:
            if word is not None:
                try:
                    word.Quit()
                except Exception:
                    pass
            pythoncom.CoUninitialize()

    out = []
    for docx_path, pdf in made:
        try:
            os.remove(docx_path)
        except OSError:
            warns.append("Đã ra PDF nhưng không xoá được .docx: %s" % docx_path)
        out.append(pdf)
    done = set(p for p, _ in made)
    out += [p for p in paths if p not in done]     # cai nao khong chuyen duoc thi giu .docx
    return out


# --------------------------------------------------------------------- chay

LOG_GOC = ("CERT | %-26s | %-2s thang | %s - %s | STBH %16s | phi %14s | %s.docx")
LOG_SDBS = ("ENDT | %-26s | don goc %-22s cap %-10s | %-5s | SL %3s | STBH %16s | "
            "phi %14s | %s.docx")


def batch(rows, tpl, filler, namer, sub_dir, merged_path, warns, log, label, progress):
    os.makedirs(sub_dir, exist_ok=True)
    docs, used = [], {}
    for i, rec in enumerate(rows, 1):
        base = safe_name(namer(rec))
        used[base] = used.get(base, 0) + 1
        if used[base] > 1:
            warns.append("%s: trùng tên file, bản thứ %d được thêm hậu tố."
                         % (base, used[base]))
            base = "%s (%d)" % (base, used[base])
        doc = filler(rec, tpl, warns)
        doc.save(os.path.join(sub_dir, base + ".docx"))
        # dung lai chinh doi tuong nay cho ban gop: file le da ghi xong ra dia,
        # merge() chi doc/chep phan tu nen khong anh huong file da luu.
        docs.append(doc)
        log.append(label(rec, base))
        if progress:
            progress(i, len(rows), base)
    merge(docs).save(merged_path)
    return merged_path


def generate(src, run_dir, want_pdf=True, progress=None, templates=TEMPLATES):
    """Sinh toan bo giay tu file Excel `src` vao thu muc `run_dir` (da ton tai).

    progress(done, total, nhan) duoc goi sau moi giay - de Streamlit ve thanh
    tien trinh. Tra ve dict: {'made', 'merged', 'warns', 'log', 'sheets'}.
    """
    tpl_gcn, tpl_sdbs = find_templates(templates)
    warns, log, merged = [], [], []
    sheets = load_input(src, set(tpl_gcn), warns)
    day = dt.datetime.now().strftime("%Y-%m-%d")

    # ngay cap don goc den tu dau - de trang web bao lai cho nguoi dung
    goc_src = {"cột Excel": 0, "dò sheet GCN gốc": 0, "": 0}
    for _n, kind, _p, rows in sheets:
        if kind == "sdbs":
            for r in rows:
                goc_src[r.get("_goc_src", "")] += 1

    total = sum(len(rows) for _n, _k, _p, rows in sheets)
    done = [0]

    def bump(_i, _n, label):
        done[0] += 1
        if progress:
            progress(done[0], total, label)

    made, info = 0, []
    for sheet, kind, prod, rows in sheets:
        if kind == "goc":
            merged.append(batch(
                rows, tpl_gcn[prod], fill_goc, lambda rec: name_goc(prod, rec),
                os.path.join(run_dir, "certificates"),
                os.path.join(run_dir, "GCNBH %s %s - merged %d certificates.docx"
                             % (prod, day, len(rows))),
                warns, log,
                lambda rec, base: LOG_GOC % (
                    txt(rec["so_don"]), thang_bh(rec, "", []), fdate(rec["tu"]),
                    fdate(rec["den"]), money(rec["stbh"]), money(rec["phi_gross"]), base),
                bump))
        else:
            if not tpl_sdbs:
                raise ToolError("Sheet '%s' là sheet SĐBS nhưng không có mẫu SĐBS "
                                "trong %s." % (sheet, templates))
            merged.append(batch(
                rows, tpl_sdbs, fill_sdbs, lambda rec: name_sdbs(prod, rec),
                os.path.join(run_dir, "endorsements"),
                os.path.join(run_dir, "SDBS %s %s - merged %d endorsements.docx"
                             % (prod, day, len(rows))),
                warns, log,
                lambda rec, base: LOG_SDBS % (
                    txt(rec["so_don"]), txt(rec["so_goc"]),
                    fdate(rec.get("ngay_cap_goc")) or "-", txt(rec["loai"]).strip(),
                    txt(rec["so_luong"]), money(rec["stbh"]), money(rec["phi_gross"]), base),
                bump))
        made += len(rows)
        info.append({"sheet": sheet, "kind": kind, "prod": prod, "rows": len(rows)})

    if want_pdf:
        merged = to_pdf(merged, warns, progress)

    with io.open(os.path.join(run_dir, "log.txt"), "w", encoding="utf-8") as f:
        f.write("Nguon  : %s\n" % os.path.basename(src))
        f.write("Chay   : %s\n\n" % dt.datetime.now().strftime("%d/%m/%Y %H:%M"))
        f.write("\n".join(log))
        f.write("\n\nCanh bao (%d):\n" % len(warns))
        f.write("\n".join(warns) if warns else "(khong co)")

    return {"made": made, "merged": merged, "warns": warns, "log": log,
            "sheets": info, "goc_src": goc_src}
