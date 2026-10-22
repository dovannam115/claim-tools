# -*- coding: utf-8 -*-
"""Reconcile TGDĐ — engine cho trang web, port tu 2 tool Excel/VBA:

  WEEKLY  : WeeklyNetPremium.bas  — file du lieu tuan (PIAS hoac PVI Connect) -> "phi sau giam"
  MONTHLY : ReconcileMonthly.bas  — cac file tuan da tinh phi x file DMX -> Results + Summary

Bang rule (CTKM, map nhom hang, rule STT, cac case doi soat) doc tu `reconcile_tgdd_rules.xlsx`
canh file nay; chay `_export_reconcile_tgdd_rules.py` de dong bo lai tu 2 file .xlsm goc.

Cac quy uoc PHAI giu dung nhu ban VBA (da doi chieu tung o):
  - phi goc + vat  = round(phi_goc_truoc_vat * VAT / 1000) * 1000   (boi so 1000)
  - phi giam + vat = round(phi_goc_cong_vat * (1 - KM))             (lam tron, khong du)
  - phi giam - vat = phi_giam_cong_vat / VAT                        (KHONG lam tron)
  - do cot theo TEN header sau khi bo dau, khong hardcode vi tri
  - ngay/thang co the la text (PVI Connect) hoac so/date (PIAS) -> tu nan ve dung kieu
"""
import os
import io
import datetime
import unicodedata

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

APP_DIR = os.path.dirname(os.path.abspath(__file__))
RULES_PATH = os.path.join(APP_DIR, "reconcile_tgdd_rules.xlsx")

SUFFIX = " - phi sau giam"

# layout file "phi sau giam": header dong 3, dong tong dong 1, data tu dong 4
OUT_COLS = [
    ("Ngày cấp đơn", 12), ("Số đơn", 22), ("Số Imei", 18), ("Người được bảo hiểm", 24),
    ("Mã sp", 9), ("MTN-GTBH (ntệ)", 14), ("phí gốc - vat", 13), ("phí giảm - vat", 14),
    ("vat", 12), ("phí giảm + vat", 14), ("Nhóm hàng", 9), ("RequestId", 34),
    ("Hãng", 14), ("tên thiết bị", 30), ("Điện thoại", 13), ("Ngày đầu", 12),
    ("Ngày cuối", 12), ("Ngày đầu BH hãng", 14), ("Ngày cuối BH hãng", 14),
    ("Tháng BH", 9), ("Tháng BH hãng", 11), ("STT", 6),
]
NCOL = len(OUT_COLS)
C_STT = 22
HL_COLS = {8, 9, 10}          # 3 cot ket qua: nen vang, chu do

BODY_FMT = {1: "dd/mm/yyyy", 2: "@", 3: "@", 5: "@", 6: "#,##0", 7: "#,##0",
            8: "#,##0.00", 9: "#,##0.00", 10: "#,##0", 12: "@", 15: "@",
            16: "dd/mm/yyyy", 17: "dd/mm/yyyy", 18: "dd/mm/yyyy", 19: "dd/mm/yyyy",
            C_STT: "0;;"}

APP_HDR = ["STT", "Kết quả", "Phí ĐMX", "Chênh lệch", "Ghi chú", "File tuần"]
APP_WIDTH = [6, 28, 13, 12, 26, 22]
NOTE = {2: "PVI giảm, ĐMX chưa giảm", 9: "lệch phí (kiểm tra tay)", 6: "có trong sheet HỦY",
        4: "ĐMX không có giao dịch này", 5: "ĐMX có, PVI không",
        "dup": "trùng RequestId trong file tuần", 10: "thiếu RequestId, không đối soát được"}
SRC_DMX = "(ĐMX)"


# ---------------------------------------------------------------- chuan hoa
def vnfold(s):
    s = str(s if s is not None else "")
    s = s.replace("Đ", "D").replace("đ", "d")
    s = unicodedata.normalize("NFD", s)
    return "".join(c for c in s if not unicodedata.combining(c)).lower()


def normh(s):
    """Chuan hoa header: bo dau + bo moi khoang trang/xuong dong."""
    t = vnfold(s)
    for ch in ("\r", "\n", "\t", " ", "\xa0"):
        t = t.replace(ch, "")
    return t


def to_num(v):
    """So that (PIAS) hoac TEXT (PVI Connect: '12', '1,234'). Dau ',' = phan cach nghin."""
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return -1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, datetime.datetime):
        return 0.0
    t = str(v).replace(",", "").replace(" ", "").replace("\xa0", "")
    try:
        return float(t)
    except ValueError:
        return 0.0


def num_or_blank(v):
    """Cot so nhung giu o RONG khi nguon rong (MTN, Tháng BH)."""
    if v is None or str(v).strip() == "":
        return None
    n = to_num(v)
    return int(n) if float(n).is_integer() else n


def to_date(v):
    """date that (PIAS) hoac TEXT 'dd/mm/yyyy[ hh:mm:ss]' (PVI Connect). Khong dung locale."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime.datetime):
        return v
    if isinstance(v, datetime.date):
        return datetime.datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)):
        if v <= 0:
            return None
        return datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(v))
    part = str(v).strip().split(" ")[0].replace("-", "/").replace(".", "/")
    p = part.split("/")
    if len(p) != 3:
        return None
    try:
        if len(p[0].strip()) == 4:
            y, m, d = int(p[0]), int(p[1]), int(p[2])
        else:
            d, m, y = int(p[0]), int(p[1]), int(p[2])
        if not (1 <= d <= 31 and 1 <= m <= 12 and 1900 <= y <= 2200):
            return None
        return datetime.datetime(y, m, d)
    except ValueError:
        return None


def as_text(v):
    """Giong CStr(Nz(v)) ben VBA — cot text khong duoc mat so 0 dau / ra scientific."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime.datetime):
        return v.strftime("%d/%m/%Y")
    return str(v)


# ---------------------------------------------------------------- bang rule
class Rules:
    def __init__(self, path=RULES_PATH):
        self.path = path
        self.mtime = os.path.getmtime(path)
        wb = openpyxl.load_workbook(path, data_only=True)
        self.vat, self.tol = 1.1, 5.0
        ws = wb["Params"]
        r = 4
        while ws.cell(r, 1).value not in (None, ""):
            k = str(ws.cell(r, 1).value).strip().lower()
            if k == "vat":
                self.vat = float(ws.cell(r, 2).value)
            elif k == "tolerance":
                self.tol = float(ws.cell(r, 2).value)
            r += 1

        self.cats = {}
        ws = wb["Categories"]
        r = 4
        while ws.cell(r, 1).value not in (None, ""):
            self.cats.setdefault(str(ws.cell(r, 1).value).strip(),
                                 str(ws.cell(r, 3).value or "").strip())
            r += 1

        self.promo = {}                     # {SANPHAM: {nhom_fold: (hang_cond, km)}}
        ws = wb["Promotions"]
        r = 4
        while ws.cell(r, 1).value not in (None, ""):
            sp = str(ws.cell(r, 1).value).strip().upper()
            nm = vnfold(str(ws.cell(r, 2).value or "").strip())
            self.promo.setdefault(sp, {}).setdefault(
                nm, (str(ws.cell(r, 3).value or "").strip(), float(ws.cell(r, 4).value or 0)))
            r += 1

        self.stt = {}                       # {SANPHAM: [rule, ...]} giu dung thu tu dong
        ws = wb["STTRules"]
        r = 4
        while ws.cell(r, 1).value not in (None, ""):
            sp = str(ws.cell(r, 1).value).strip().upper()
            bh = str(ws.cell(r, 6).value if ws.cell(r, 6).value is not None else "").strip()
            bhh = str(ws.cell(r, 7).value if ws.cell(r, 7).value is not None else "").strip()
            self.stt.setdefault(sp, []).append({
                "stt": int(to_num(ws.cell(r, 2).value)),
                "masp": str(ws.cell(r, 3).value or "").strip(),
                "toks": [t.strip() for t in str(ws.cell(r, 4).value or "").split(",")],
                "hang": vnfold(str(ws.cell(r, 5).value or "").strip()),
                "bh": 0 if bh in ("", "*") else int(to_num(bh)),
                "bhh": 0 if bhh in ("", "*") else int(to_num(bhh)),
            })
            r += 1

        self.nhommap = {}                   # {SANPHAM: [(goc_fold, kw_fold, chuan), ...]}
        ws = wb["NhomMap"]
        r = 4
        while ws.cell(r, 1).value not in (None, ""):
            sp = str(ws.cell(r, 1).value).strip().upper()
            self.nhommap.setdefault(sp, []).append(
                (vnfold(str(ws.cell(r, 2).value or "").strip()),
                 vnfold(str(ws.cell(r, 3).value or "").strip()),
                 str(ws.cell(r, 4).value or "").strip()))
            r += 1

        self.cases = []                     # 10 dong mo ta case doi soat thang
        ws = wb["MonthlyCases"]
        r = 4
        while ws.cell(r, 1).value not in (None, ""):
            self.cases.append([ws.cell(r, c).value for c in range(1, 6)])
            r += 1
        wb.close()

    @property
    def products(self):
        return sorted(set(list(self.promo) + list(self.stt)))

    # --- KM theo nhom hang + hang
    def km(self, product, ctkm_name, hang):
        if not ctkm_name:
            return 0.0
        table = self.promo.get(product.upper(), {})
        nm = vnfold(str(ctkm_name).strip())
        if nm not in table:
            return 0.0
        cond, km = table[nm]
        c, h = vnfold(cond), vnfold(str(hang or "").strip())
        if "tru " in c or "(tru" in c:
            if h and h in c:
                km = 0.0
        elif "tat ca" not in c and c:
            if not h or h not in c:
                km = 0.0
        return km

    # --- nan ten nhom theo ten thiet bi (chi sp co dong trong NhomMap, vd MDM)
    def apply_nhommap(self, product, nhom_name, ten_tb):
        rows = self.nhommap.get(product.upper(), [])
        if not rows:
            return nhom_name
        ten, goc = vnfold(ten_tb), vnfold(str(nhom_name).strip())
        for g, kw, chuan in rows:                       # pass 1: dong co tu khoa
            if g == goc and kw:
                if any(k.strip() and k.strip() in ten for k in kw.split("|")):
                    return chuan
        for g, kw, chuan in rows:                       # pass 2: dong "con lai"
            if g == goc and not kw:
                return chuan
        return nhom_name

    # --- rule STT: khop dau tien thang
    def match_stt(self, product, masp, nhom, hang, bh, bhh):
        h = vnfold(str(hang or "").strip())
        for ru in self.stt.get(product.upper(), []):
            if ru["masp"] not in ("", "*") and str(masp).strip() != ru["masp"]:
                continue
            ok = False
            for t in ru["toks"]:
                if t in ("*", ""):
                    ok = True
                elif t.endswith("*"):
                    pre = t[:-1]
                    if pre == "" or str(nhom).startswith(pre):
                        ok = True
                elif t == str(nhom):
                    ok = True
                if ok:
                    break
            if not ok:
                continue
            c = ru["hang"]
            if c in ("", "*") or "tat ca hang" in c or "tat ca cac hang" in c:
                pass
            elif "tru " in c or "(tru" in c:
                if h and h in c:
                    continue
            else:
                if not h or h not in c:
                    continue
            if ru["bh"] > 0 and bh != ru["bh"]:
                continue
            if ru["bhh"] > 0 and bhh != ru["bhh"]:
                continue
            return ru["stt"]
        return 0


# ------------------------------------------------------------ doc file nguon
def _rewind(src):
    """File upload cua Streamlit doc lai lan 2 se rong neu khong tua ve dau."""
    if hasattr(src, "seek"):
        try:
            src.seek(0)
        except Exception:
            pass
    return src


def _rows_of(src, sheet_pick="requestid"):
    """Doc sheet co cot moc (vd RequestId) cua 1 file Excel -> list cac tuple."""
    wb = openpyxl.load_workbook(_rewind(src), read_only=True, data_only=True)
    try:
        chosen, rows = None, None
        for name in wb.sheetnames:
            data = list(wb[name].iter_rows(values_only=True))
            if _find_hdr(data, sheet_pick, 8) is not None:
                chosen, rows = name, data
                break
        if chosen is None:
            chosen = wb.sheetnames[0]
            rows = list(wb[chosen].iter_rows(values_only=True))
        return chosen, rows
    finally:
        wb.close()


def _find_hdr(rows, token, max_scan):
    for i, r in enumerate(rows[:max_scan]):
        for v in r:
            if v is not None and token in normh(v):
                return i
    return None


def _pick(hdr, name, exact, dflt):
    """Vi tri cot (1-based) theo ten header; khong thay thi dung vi tri du phong."""
    for j, v in enumerate(hdr, 1):
        h = normh(v)
        if (h == name) if exact else (name in h):
            return j
    return dflt


def _cell(row, col):
    return row[col - 1] if 0 < col <= len(row) else None


# ------------------------------------------------------------------- WEEKLY
def calc_weekly(src, product="BHMR", rules=None, progress=None):
    """src = duong dan / file-like file du lieu tuan. Tra ve (rows, info)."""
    rules = rules or Rules()
    sheet, data = _rows_of(src, "requestid")
    hr = _find_hdr(data, "requestid", 8)
    if hr is None:
        raise ValueError("Không tìm thấy cột RequestId trong file, kiểm tra lại file xuất.")
    hdr = data[hr]
    col = {
        "req": _pick(hdr, "requestid", True, 15), "ngaycap": _pick(hdr, "ngaycapdon", True, 3),
        "sodon": _pick(hdr, "sodon", True, 4), "imei": _pick(hdr, "soimei", True, 5),
        "ndbh": _pick(hdr, "nguoiduocbaohiem", True, 6), "masp": _pick(hdr, "masp", True, 8),
        "mtn": _pick(hdr, "mtn", False, 9), "phi": _pick(hdr, "sotienphi", True, 11),
        "thangbh": _pick(hdr, "thangbh", True, 12), "nhom": _pick(hdr, "nhomhang", True, 13),
        "hang": _pick(hdr, "hang", True, 17), "ten": _pick(hdr, "tenthietbi", True, 18),
        "dt": _pick(hdr, "dienthoai", True, 19), "ngaydau": _pick(hdr, "ngaydau", True, 20),
        "ngaycuoi": _pick(hdr, "ngaycuoi", True, 21), "bhdau": _pick(hdr, "ngaydaubhhang", True, 22),
        "bhcuoi": _pick(hdr, "ngaycuoibhhang", True, 23),
        "thangbhh": _pick(hdr, "thangbhhang", True, 24),
    }
    vat = rules.vat
    has_nhommap = bool(rules.nhommap.get(product.upper()))
    out, total = [], len(data)
    for i, row in enumerate(data):
        if progress and i % 2000 == 0:
            progress(i, total)
        req = as_text(_cell(row, col["req"])).strip()
        if not req or req.upper() == "REQUESTID":       # bo header + dong chen "ID: ..."
            continue
        fee = to_num(_cell(row, col["phi"]))
        code = as_text(_cell(row, col["nhom"])).strip()
        hang = as_text(_cell(row, col["hang"])).strip()
        km = rules.km(product, rules.cats.get(code, ""), hang)
        goc_plus = round(fee * vat / 1000) * 1000
        giam_plus = round(goc_plus * (1 - km))
        giam_minus = giam_plus / vat
        bh = int(to_num(_cell(row, col["thangbh"])))
        bhh = int(to_num(_cell(row, col["thangbhh"])))
        masp = as_text(_cell(row, col["masp"])).strip()
        if has_nhommap:
            nh_name = rules.cats.get(code) or code
            nhom_key = rules.apply_nhommap(product, nh_name,
                                           as_text(_cell(row, col["ten"])).strip())
        else:
            nhom_key = code
        out.append([
            to_date(_cell(row, col["ngaycap"])),
            as_text(_cell(row, col["sodon"])),
            as_text(_cell(row, col["imei"])),
            _cell(row, col["ndbh"]),
            as_text(_cell(row, col["masp"])),
            num_or_blank(_cell(row, col["mtn"])),
            fee, giam_minus, giam_plus - giam_minus, giam_plus,
            _cell(row, col["nhom"]),
            req,
            _cell(row, col["hang"]),
            _cell(row, col["ten"]),
            as_text(_cell(row, col["dt"])),
            to_date(_cell(row, col["ngaydau"])),
            to_date(_cell(row, col["ngaycuoi"])),
            to_date(_cell(row, col["bhdau"])),
            to_date(_cell(row, col["bhcuoi"])),
            num_or_blank(_cell(row, col["thangbh"])),
            num_or_blank(_cell(row, col["thangbhh"])),
            rules.match_stt(product, masp, nhom_key, hang, bh, bhh),
        ])
    if progress:
        progress(total, total)
    info = {"sheet": sheet, "header_row": hr + 1, "rows": len(out),
            "unmatched_stt": sum(1 for r in out if not r[C_STT - 1]),
            "sum_fee": sum(r[6] for r in out), "sum_giam_plus": sum(r[9] for r in out),
            "sum_giam_minus": sum(r[7] for r in out), "sum_vat": sum(r[8] for r in out)}
    return out, info


def write_weekly_xlsx(rows):
    """Dung dung layout file sach cua tool VBA (Sheet1, header dong 3, tong dong 1)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.sheet_view.showGridLines = False

    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hdr_plain = Font(name="Arial", bold=True, size=10, color="000000")
    hdr_hl = Font(name="Arial", bold=True, size=10, color="FF0000")
    fill_yellow = PatternFill("solid", fgColor="FFFF00")
    fill_tot = PatternFill("solid", fgColor="DCE6F1")
    bold = Font(name="Arial", bold=True, size=10)

    for j, (h, w) in enumerate(OUT_COLS, 1):
        c = ws.cell(3, j, h)
        c.font = hdr_hl if j in HL_COLS else hdr_plain
        if j in HL_COLS:
            c.fill = fill_yellow
        c.alignment = center
        c.border = border
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[3].height = 26

    for i, row in enumerate(rows, 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            fmt = BODY_FMT.get(j)
            if fmt:
                c.number_format = fmt

    last = len(rows) + 3
    ws.row_dimensions[1].height = 18
    c = ws.cell(1, 2, "TỔNG")
    c.font = Font(name="Arial", bold=True, size=10, color="183665")
    c.alignment = Alignment(horizontal="right", vertical="center")
    if rows:
        ws.cell(1, 5, "=COUNTA(L4:L%d)" % last).number_format = "#,##0"
        for j, col in ((6, "F"), (7, "G")):
            ws.cell(1, j, "=SUM(%s4:%s%d)" % (col, col, last)).number_format = "#,##0"
        for j, col in ((8, "H"), (9, "I")):
            ws.cell(1, j, "=SUM(%s4:%s%d)" % (col, col, last)).number_format = "#,##0.00"
        ws.cell(1, 10, "=SUM(J4:J%d)" % last).number_format = "#,##0"
        ws.cell(1, 16, "=MIN(P4:P%d)" % last).number_format = "dd/mm/yyyy"
        ws.cell(1, 17, "=MAX(Q4:Q%d)" % last).number_format = "dd/mm/yyyy"
        for j in (5, 6, 7, 8, 9, 10, 16, 17):
            ws.cell(1, j).font = bold
            ws.cell(1, j).fill = fill_tot
            ws.cell(1, j).alignment = center
        # STT khong khop rule (= 0) -> to do, giong conditional format cua tool
        ws.conditional_formatting.add(
            "%s4:%s%d" % (get_column_letter(C_STT), get_column_letter(C_STT), last),
            CellIsRule(operator="equal", formula=["0"],
                       fill=PatternFill("solid", fgColor="FFC7CE"),
                       font=Font(color="9C0006")))
        ws.auto_filter.ref = "A3:%s%d" % (get_column_letter(NCOL), last)
    ws.freeze_panes = "A4"
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def weekly_out_name(src_name):
    base = os.path.splitext(os.path.basename(src_name))[0]
    return base + SUFFIX + ".xlsx"


# ------------------------------------------------------------------ MONTHLY
def _dmx_dict(rows, mg_tok, amt_tok, amt_tok2=""):
    hr = _find_hdr(rows, mg_tok, 10)
    if hr is None:
        hr = 0
    hdr = rows[hr]
    mgc = _pick(hdr, mg_tok, False, 0)
    amtc = _pick(hdr, amt_tok, False, 0)
    if not amtc and amt_tok2:
        amtc = _pick(hdr, amt_tok2, False, 0)
    if not mgc:
        raise ValueError("Không tìm thấy cột [%s] trong file ĐMX." % mg_tok)
    if not amtc:
        raise ValueError("Không tìm thấy cột [%s] trong file ĐMX." % amt_tok)
    d = {}
    for row in rows[hr + 1:]:
        mg = _cell(row, mgc)
        if isinstance(mg, str) and len(mg.strip()) >= 24:
            d.setdefault(mg.strip(), to_num(_cell(row, amtc)))
    return d


def _row_all_blank(row, rc, gc, fc):
    for c in (1, rc, gc, fc):
        if str(_cell(row, c) if _cell(row, c) is not None else "").strip():
            return False
    return True


def reconcile_monthly(weekly, dmx, rules=None, progress=None):
    """weekly = [(ten_file, path/bytes), ...] file tuan da tinh phi; dmx = file DMX.
    Tra ve (header, rows, counts, sums)."""
    rules = rules or Rules()
    tol = rules.tol

    files = []
    for name, src in weekly:
        _, data = _rows_of(src, "requestid")
        files.append((name, data))
    if not files:
        raise ValueError("Chưa có file tuần nào.")

    a0 = files[0][1]
    hr0 = _find_hdr(a0, "requestid", 10)
    hr0 = 2 if hr0 is None else hr0
    wk_ncol = max(len(r) for r in a0)
    base_hdr = [as_text(v) for v in a0[hr0]] + [""] * (wk_ncol - len(a0[hr0]))
    req_col = _pick(base_hdr, "requestid", False, 12)

    # --- DMX: sheet BAN / HUY do theo ten sheet
    wb = openpyxl.load_workbook(_rewind(dmx), read_only=True, data_only=True)
    try:
        sb = sh = None
        for name in wb.sheetnames:
            f = vnfold(name)
            if sb is None and "ban" in f:
                sb = name
            if sh is None and "huy" in f:
                sh = name
        if not sb or not sh:
            raise ValueError("File ĐMX không có sheet BÁN / HỦY.")
        ban = _dmx_dict(list(wb[sb].iter_rows(values_only=True)), "magiaodich", "thanhtien")
        huy = _dmx_dict(list(wb[sh].iter_rows(values_only=True)), "magiaodich",
                        "sotiendongy", "dongy")
    finally:
        wb.close()

    counts = {i: 0 for i in range(1, 11)}
    sums = {i: 0.0 for i in range(1, 11)}
    seen = {}
    out = []
    total = sum(len(d) for _, d in files) + len(ban)
    step = 0
    for name, data in files:
        hr = _find_hdr(data, "requestid", 10)
        hr = 2 if hr is None else hr
        hdr = [as_text(v) for v in data[hr]]
        rc = _pick(hdr, "requestid", False, 12)
        gc = _pick(hdr, "phigiam+vat", False, 10)
        fc = _pick(hdr, "phigoc-vat", False, 7)
        for row in data[hr + 1:]:
            step += 1
            if progress and step % 2000 == 0:
                progress(step, total)
            if _row_all_blank(row, rc, gc, fc):
                continue
            vals = [_cell(row, c) for c in range(1, wk_ncol + 1)]
            req = _cell(row, rc)
            key = req.strip() if isinstance(req, str) and len(req.strip()) >= 24 else ""
            giam = to_num(_cell(row, gc))
            fee = to_num(_cell(row, fc))
            phidmx, note, has = None, "", False
            if not key:
                cs = 10
                note = NOTE[10]
            elif key not in seen:
                seen[key] = giam
                if key in ban:
                    b = to_num(ban[key])
                    goc = round(fee * 1.1 / 1000) * 1000
                    if abs(giam - b) <= tol:
                        cs = 1
                    elif goc != giam and abs(goc - b) <= tol:
                        cs, note = 2, NOTE[2]
                    else:
                        cs, note = 9, NOTE[9]
                    has, phidmx = True, b
                elif key in huy:
                    cs, note = 6, NOTE[6]
                    has, phidmx = True, to_num(huy[key])
                else:
                    cs, note = 4, NOTE[4]
            else:
                cs = 7 if abs(to_num(seen[key]) - giam) <= tol else 8
                note = NOTE["dup"]
            out.append(vals + [cs, rules.cases[cs - 1][4],
                               phidmx if has else None,
                               (giam - phidmx) if has else None, note, name])
            counts[cs] += 1
            sums[cs] += giam

    # --- DMX co, PVI khong (case 5)
    for k, v in ban.items():
        if k not in seen:
            step += 1
            if progress and step % 2000 == 0:
                progress(step, total)
            vals = [None] * wk_ncol
            if 0 < req_col <= wk_ncol:
                vals[req_col - 1] = k
            out.append(vals + [5, rules.cases[4][4], to_num(v), None, NOTE[5], SRC_DMX])
            counts[5] += 1
            sums[5] += to_num(v)
    if progress:
        progress(total, total)
    return base_hdr, out, counts, sums


def _fmt_for_header(h):
    nh = normh(h)
    if any(t in nh for t in ("imei", "sodon", "masp", "requestid", "dienthoai")):
        return "@"
    if nh.startswith("ngay"):
        return "dd/mm/yyyy"
    if "phi" in nh or nh == "vat" or "mtn" in nh:
        return "#,##0"
    return None


def write_monthly_xlsx(base_hdr, rows, counts, sums, rules=None):
    rules = rules or Rules()
    wk_ncol = len(base_hdr)
    total_cols = wk_ncol + len(APP_HDR)
    stt_c, kq_c, phi_c, ch_c, note_c, file_c = (wk_ncol + i for i in range(1, 7))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    navy, green, purple = "183665", "1E7145", "6A2C91"
    hdr_font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for j, h in enumerate(base_hdr + APP_HDR, 1):
        c = ws.cell(1, j, h)
        c.font = hdr_font
        c.alignment = center
        c.fill = PatternFill("solid",
                             fgColor=navy if j <= wk_ncol else (green if j == phi_c else purple))
    ws.row_dimensions[1].height = 28

    for i, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v)
    last = len(rows) + 1
    if rows:
        for j, h in enumerate(base_hdr, 1):
            fmt = _fmt_for_header(h)
            if fmt:
                for i in range(2, last + 1):
                    ws.cell(i, j).number_format = fmt
        for i in range(2, last + 1):
            ws.cell(i, stt_c).number_format = "0"
            ws.cell(i, phi_c).number_format = "#,##0"
            ws.cell(i, ch_c).number_format = "#,##0"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(total_cols), max(last, 1))
    ws.freeze_panes = "A2"
    for j, w in zip((stt_c, kq_c, phi_c, ch_c, note_c, file_c), APP_WIDTH):
        ws.column_dimensions[get_column_letter(j)].width = w

    # ---- Summary
    sm = wb.create_sheet("Summary")
    sm.sheet_view.showGridLines = False
    c = sm.cell(1, 1, "TỔNG HỢP KẾT QUẢ ĐỐI SOÁT")
    c.font = Font(name="Arial", bold=True, size=13, color=navy)
    head = ["STT", "Đối soát", "Trường hợp", "Ghi chú", "Kết quả", "Số dòng",
            "Tổng phí giảm + vat"]
    for j, h in enumerate(head, 1):
        cc = sm.cell(3, j, h)
        cc.font = hdr_font
        cc.fill = PatternFill("solid", fgColor=navy)
        cc.alignment = center
    for i, case in enumerate(rules.cases, 4):
        n = i - 3
        for j, v in enumerate(case, 1):
            sm.cell(i, j, v).font = Font(name="Arial", size=10)
            sm.cell(i, j).alignment = Alignment(vertical="center", wrap_text=True)
        sm.cell(i, 6, counts.get(n, 0)).number_format = "#,##0"
        sm.cell(i, 7, sums.get(n, 0.0)).number_format = "#,##0"
    r = 4 + len(rules.cases)
    sm.cell(r, 5, "TỔNG").font = Font(name="Arial", bold=True, size=10)
    sm.cell(r, 6, "=SUM(F4:F%d)" % (r - 1)).number_format = "#,##0"
    sm.cell(r, 7, "=SUM(G4:G%d)" % (r - 1)).number_format = "#,##0"
    for j in (6, 7):
        sm.cell(r, j).font = Font(name="Arial", bold=True, size=10)
    for col, w in zip("ABCDEFG", (6, 32, 46, 40, 40, 12, 20)):
        sm.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()
