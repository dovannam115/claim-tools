# -*- coding: utf-8 -*-
"""Search Database — tra cuu THANG vao bang phi tren Fabric (dbo.D_DATA_PREMIUM_EQUIPMENT_V2).

Khac Map Claim: KHONG loc, KHONG cham diem, KHONG chon "don dung".
   Nhap IMEI / SDT / so hop dong / ma giao dich -> tra ve MOI dong phi khop khoa, du
   ca dong huy (Quantity = -1), ca cac ban sua doi /E. Nguoi dung tu doc.
   Dung khi Map Claim bao NO MATCH hoac khi can biet mot thiet bi co nhung don nao.

Cot tra ve = TOAN BO cot cua bang phi (SELECT *) + [Nhom san pham] (cung CASE voi
   vw_Claim_Premium_Matched) + co cau tai tu dbo.D_RI_STRUCTURE (tong share theo HOP DONG
   + share tung nha tai) — giong het bo cot ma Map Claim xuat ra.

KHONG sua map_claim_ri.py: module do da duoc Streamlit import san, sua la phai restart
   ca service dung chung cua phong Claims. Trang nay chi DOC cac hang so cua no.
"""
import io, datetime, decimal, traceback

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import map_claim_ri as core

core.HEADLESS = True          # die() phai raise MapError, khong duoc print + input()
core.GUI = False              # khong bao gio bung dialog PowerShell tren server

MAX_VALS = 500                # tran so gia tri 1 lan tra, tranh quet nhieu lan bang 4 trieu dong
# Tran so DONG tra ve. KHONG phai lam canh: da gap that khi test tren warehouse —
#   · so hop dong GOP ca chuong trinh: 1 so don 1D1_2026 tra 12.367 dong phi;
#   · gia tri RAC trong database: co 21.232 dong ghi [RequestID] = 'x'.
# Khong chan thi go 1 gia tri cung treo trinh duyet. Nguoi dung tu nang duoc trong UI.
MAX_ROWS = 5000
IN_CHUNK = 200                # giong fetch_in cua engine
LIKE_CHUNK = 50               # so hop dong dung LIKE tien to -> OR it hon
FAM_COL = "Nhóm sản phẩm"     # alias cua FAM_SQL; dat tieng Viet de khong dung ten cot that
RI_TOT_COL = "RI Share (hợp đồng)"

KEY_LABELS = {
    "serial": "IMEI / Serial",
    "req":    "Mã giao dịch",
    "phone":  "Số điện thoại",
    "policy": "Số hợp đồng",
}
KEY_EXPR = {                                   # y het khoa ma match() dung
    "serial": "LTRIM(RTRIM([IMEI]))",
    "req":    "LTRIM(RTRIM([RequestID]))",
    "phone":  "RIGHT(LTRIM(RTRIM([Phone])),9)",
}
# Cot dua len dau bang cho de doc; cac cot con lai cua bang phi xep sau, giu thu tu database.
MAIN_COLS = ["Khoá tra", "Giá trị nhập", "IMEI", "Phone", "RequestID", "Policy No",
             "DataSource", "Program", "Program Coverage", "Effective Date", "Expired Date",
             "SI", "Quantity", "Premium Gross", "RI Share Pct", RI_TOT_COL] \
            + core.REINSURERS + [FAM_COL]


# ------------------------------------------------------------------ nhap lieu
def tach_gia_tri(text):
    """Tach o nhap thanh danh sach gia tri: moi dong 1 gia tri, cho phep them , ; tab."""
    out, seen = [], set()
    for chunk in text.replace(";", "\n").replace(",", "\n").replace("\t", "\n").split("\n"):
        v = core.stext(chunk)
        if not v:
            continue
        if v not in seen:
            seen.add(v)
            out.append(v)
    return out


def _esc_like(s):
    """Chan ky tu dai dien cua LIKE lot tu o nhap ([ % _) — so don chi co chu, so va '/'."""
    return s.replace("[", "[[]").replace("%", "[%]").replace("_", "[_]")


# ------------------------------------------------------------------ truy van
def _rows(cur, where, params, top):
    """SELECT * + Nhom san pham. Lay het cot cua bang phi theo dung ten database.
       top: chan cung so dong 1 cau SQL keo ve (khong ORDER BY -> lay dong nao la tuy DB)."""
    sel = (f"SELECT TOP {int(top)} *, ({core.FAM_SQL}) AS [{FAM_COL}] "
           f"FROM {core.SRCTBL} WHERE {where}")
    cur.execute(sel, *params)
    cols = [d[0] for d in cur.description]
    return [dict(zip(cols, r)) for r in cur.fetchall()]


def _fetch_key(cur, kind, values, top):
    """Tra theo 1 khoa. -> ({gia tri nhap: [dong phi]}, cham_tran)
       cham_tran = mot cau SQL tra ve dung `top` dong => database CON dong nua, da bi cat."""
    hit, capped = {}, False
    if kind == "policy":
        # So hop dong go tay thuong THIEU duoi /E01 -> khop TIEN TO, khong khop tuyet doi.
        for i in range(0, len(values), LIKE_CHUNK):
            chunk = values[i:i + LIKE_CHUNK]
            where = " OR ".join(["UPPER(LTRIM(RTRIM([Policy No]))) LIKE ?"] * len(chunk))
            rows = _rows(cur, where, [_esc_like(v.upper()) + "%" for v in chunk], top)
            capped = capped or len(rows) >= top
            for r in rows:
                pol = (core.stext(r.get("Policy No")) or "").upper()
                for v in chunk:
                    if pol.startswith(v.upper()):
                        hit.setdefault(v, []).append(r)
        return hit, capped

    expr, want = KEY_EXPR[kind], values
    if kind == "phone":
        want = [core.last9(v) for v in values]
    pairs = [(v, w) for v, w in zip(values, want) if w]
    if not pairs:
        return hit, capped
    uniq = sorted({w for _, w in pairs})
    got = {}
    for i in range(0, len(uniq), IN_CHUNK):
        chunk = uniq[i:i + IN_CHUNK]
        ph = ",".join("?" * len(chunk))
        rows = _rows(cur, f"{expr} IN ({ph})", chunk, top)
        capped = capped or len(rows) >= top
        for r in rows:
            if kind == "serial":
                k = (core.stext(r.get("IMEI")) or "").strip()
            elif kind == "req":
                k = (core.stext(r.get("RequestID")) or "").strip()
            else:
                k = core.last9(r.get("Phone"))
            if k:
                got.setdefault(k, []).append(r)
    for v, w in pairs:
        if w in got:
            hit.setdefault(v, []).extend(got[w])
    return hit, capped


def tra_cuu(cur, mode, values, top=MAX_ROWS):
    """-> ([(gia tri nhap, ten khoa, dong phi)], bi_cat)
       Moi GIA TRI NHAP la mot luot tra doc lap. mode='auto': gia tri co '/' tra theo so hop
       dong; con lai di dung thu tu uu tien cua engine IMEI -> Ma giao dich -> SDT, trung khoa
       nao thi dung o khoa do (y het match())."""
    out, capped = [], False
    # dict(r): MOT dong phi co the khop NHIEU gia tri nhap (vd nhap ca so don goc lan /E01,
    # hoac 2 SDT cung 9 so cuoi). Dung chung 1 dict thi cot "Gia tri nhap"/"Khoa tra" gan sau
    # se de len cai truoc -> moi dong deu hien gia tri cuoi cung.
    if mode != "auto":
        hit, capped = _fetch_key(cur, mode, values, top)
        for v, rows in hit.items():
            out += [(v, KEY_LABELS[mode], dict(r)) for r in rows]
    else:
        pol = [v for v in values if "/" in v]
        rest = [v for v in values if "/" not in v]
        # Hai nhom TACH ROI (co '/' hay khong), nen thu tu giua chung khong doi ket qua cua
        # tung gia tri — nhung khi CHAM TRAN thi co: mot so hop dong gop keo 12.000 dong se
        # an het cho cua cac IMEI go cung o. Nen chay khoa HEP truoc, so hop dong sau cung.
        for kind, vals in (("serial", rest), ("req", rest), ("phone", rest), ("policy", pol)):
            done = {v for v, _, _ in out}
            vals = [v for v in vals if v not in done]
            if not vals:
                continue
            rem = top - len(out)
            if rem <= 0:
                capped = True
                continue
            hit, cap = _fetch_key(cur, kind, vals, rem)
            capped = capped or cap
            for v, rows in hit.items():
                out += [(v, KEY_LABELS[kind], dict(r)) for r in rows]
    return out[:top], bool(capped or len(out) > top)


def gan_co_cau_tai(r):
    """Them tong share theo HOP DONG + share tung nha tai (dbo.D_RI_STRUCTURE).
       Bang chua nap -> cac cot de trong, khong lam hong gi."""
    k = core._ri_key(r.get("DataSource"), r.get("Program Coverage"), r.get("RI Share Pct"))
    r[RI_TOT_COL] = core.RI_TOTAL.get(k)
    share = core.RI_STRUCT.get(k, {})
    for n in core.REINSURERS:
        r[n] = share.get(n)
    return r


def sap_cot(rows):
    """Thu tu cot: cot chinh truoc, phan con lai giu nguyen thu tu cua database."""
    seen = []
    for r in rows:
        for c in r:
            if c not in seen:
                seen.append(c)
    return [c for c in MAIN_COLS if c in seen] + [c for c in seen if c not in MAIN_COLS]


# ------------------------------------------------------------------ hien thi
def _disp(v):
    """Ep ve kieu Streamlit ve duoc: cot lan lon kieu se lam pyarrow no."""
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return core.fmt_d(v.date() if isinstance(v, datetime.datetime) else v)
    if isinstance(v, decimal.Decimal):
        return f"{float(v):,.0f}" if float(v).is_integer() else f"{float(v):,.2f}"
    if isinstance(v, float):
        return f"{v:,.0f}" if v.is_integer() else f"{v:,.2f}"
    if isinstance(v, int) and not isinstance(v, bool):
        return f"{v:,}"
    if isinstance(v, (bytes, bytearray)):
        return v.hex()
    return str(v)


PCT_COLS = set(core.REINSURERS) | {"RI Share Pct", RI_TOT_COL}
# Dinh danh: LUON hien dang chu. Excel/pyodbc tra IMEI ve kieu so -> format nghin se ra
# "351,954,790,151,594", con float thi ra 3.51955E+14 (dung bay stext cua engine).
ID_COLS = {"IMEI", "Phone", "RequestID", "Policy No", "Giá trị nhập"}


def _cell(col, v):
    if col in PCT_COLS:
        return core.fmt_pct(core.pnum(v))
    if col in ID_COLS:
        return core.stext(v) or ""
    return _disp(v)


def bang_hien_thi(recs, cols):
    return [{c: _cell(c, r.get(c)) for c in cols} for _v, _k, r in recs]


# ------------------------------------------------------------------ xuat Excel
def xuat_excel(recs, cols):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tra cuu"
    box = Border(*[Side(style="thin")] * 4)
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        c.fill = PatternFill("solid", start_color=core.NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box
        w = 28 if h in ("Program", "Program Coverage", "Policy No", "IMEI") else len(str(h)) + 5
        ws.column_dimensions[get_column_letter(i)].width = max(11, min(34, w))
    for r_i, (_v, _k, r) in enumerate(recs, 2):
        for i, h in enumerate(cols, 1):
            v = r.get(h)
            if h in PCT_COLS:
                v = core.pnum(v)
            elif isinstance(v, (bytes, bytearray)):
                v = v.hex()
            elif isinstance(v, decimal.Decimal):
                v = float(v)
            elif isinstance(v, datetime.datetime):
                v = v.date()
            c = ws.cell(row=r_i, column=i, value=v)
            c.font = Font(name="Arial", size=10)
            c.border = box
            c.alignment = Alignment(vertical="top")
            if h in PCT_COLS:
                c.number_format = "0.0%"
            elif isinstance(v, datetime.date):
                c.number_format = "dd/mm/yyyy"
            elif h in ID_COLS:
                # Dinh danh phai la TEXT, khong Excel hien 3.51955E+14 (bay stext cua engine).
                c.value = core.stext(v) if v is not None else None
                c.number_format = "@"
            elif isinstance(v, (int, float)) and h in ("SI", "Premium Gross", "Quantity"):
                c.number_format = "#,##0"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(recs) + 1}"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


