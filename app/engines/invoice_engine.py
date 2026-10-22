# -*- coding: utf-8 -*-
"""invoice_engine — lõi tra cứu & đối soát hóa đơn điện tử TGDĐ/ĐMX.

Dựng lại từ engine.py trong Invoice.exe (bản desktop 22/06/2026) để dùng chung cho
web Claim & Equipment Tools. Pure HTTP + ddddocr captcha — KHÔNG dùng trình duyệt, nhẹ CPU.

Luồng: getcaptchaimage -> ddddocr đọc captcha -> POST Home/ListInvoice (SĐT + số HĐ)
-> lấy link "TẢI XML" -> parse XML -> so IMEI & giá -> KHỚP / LỆCH / KHÔNG TRA ĐƯỢC.
"""
import os, sys

# ddddocr/onnxruntime mặc định ăn hết core; ghim 1 luồng để không giành CPU với các app khác
for _v in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS",
           "NUMEXPR_NUM_THREADS", "ORT_NUM_THREADS"):
    os.environ.setdefault(_v, "1")

import re, csv, time, random, threading
import requests
from xml.etree import ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

HOME = "https://hddt.thegioididong.com/"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")

# tốc độ = số luồng chạy song song (bản exe: Light 3 / Balanced 6 / Fast 10)
MODES = {"Light": 3, "Balanced": 6, "Fast": 10}

ocr = None
ocr_lock = threading.Lock()


def _ensure_ocr():
    """Nạp ddddocr một lần, ghim onnxruntime về 1 luồng."""
    global ocr
    if ocr is not None:
        return
    try:
        import onnxruntime as _ort
        _ort.set_default_logger_severity(3)
    except Exception:
        pass
    try:
        _orig = _ort.InferenceSession.__init__

        def _patched(self, *a, **k):
            so = k.get("sess_options") or _ort.SessionOptions()
            so.intra_op_num_threads = 1
            so.inter_op_num_threads = 1
            k["sess_options"] = so
            return _orig(self, *a, **k)

        _ort.InferenceSession.__init__ = _patched
    except Exception:
        pass
    import ddddocr
    ocr = ddddocr.DdddOcr(show_ad=False)


def solve(img):
    _ensure_ocr()
    with ocr_lock:                       # model không thread-safe
        return re.sub("[^A-Z0-9]", "", (ocr.classification(img) or "").upper())


def norm(s):
    return re.sub("[^A-Za-z0-9]", "", s or "").upper()


# --------------------------------------------------------------------- input
def read_input(path):
    """Đọc file danh sách máy. CSV: đúng 5 cột imei/phone/sohd/gia_khai/model.
    Excel: tự dò dòng tiêu đề (30 dòng đầu, mọi sheet) theo tên cột tiếng Việt."""
    rows = []
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                rows.append({"imei": r.get("imei", ""), "phone": r.get("phone", ""),
                             "sohd": r.get("sohd", ""), "gia_khai": r.get("gia_khai", ""),
                             "model": r.get("model", "")})
        return rows

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    def hkey(v):
        return re.sub(r"\s+", " ", str(v or "")).strip().lower()

    for ws in wb.worksheets:
        hdr_row = None
        cols = {}
        for rr in range(1, min(31, ws.max_row + 1)):
            vals = [hkey(c.value) for c in ws[rr]]
            ci = {}
            for idx, h in enumerate(vals, start=1):
                if not h:
                    continue
                if "imei" in h or "serial" in h:
                    ci.setdefault("imei", idx)
                if ("giá" in h and ("mua" in h or "khách" in h)) or "gia_khai" in h:
                    ci.setdefault("gia", idx)
                if ("kiểu" in h and "thiết" in h) or h == "model" or "tên sản" in h or "kieu thiet" in h:
                    ci.setdefault("model", idx)
                if "số hóa đơn" in h or "so hoa don" in h or h == "sohd" or "số hđ" in h:
                    ci.setdefault("sohd", idx)
                if h in ("đt", "sđt", "phone") or "điện thoại" in h:
                    ci.setdefault("phone", idx)
            if "imei" in ci:
                hdr_row, cols = rr, ci
                break
        if hdr_row is None:
            continue
        for rr in range(hdr_row + 1, ws.max_row + 1):
            row = ws[rr]

            def cell(k):
                i = cols.get(k)
                if i and i - 1 < len(row):
                    return row[i - 1].value
                return None

            imei = str(cell("imei") or "").strip()
            if not imei:
                continue
            gia = str(cell("gia") or "").strip()
            model = str(cell("model") or "").strip()
            sohd = str(cell("sohd") or "").strip()
            phone = str(cell("phone") or "").strip()
            rowtext = " ".join(str(c.value) for c in row if c.value is not None)
            if not sohd:                       # vớt số HĐ nằm lẫn trong text
                m = re.search(r"H[ĐD]\s*[:：]?\s*(\d{4,})", rowtext)
                sohd = m.group(1) if m else ""
            if not phone:
                m = re.search(r"(0\d{8,10})", rowtext)
                phone = m.group(1) if m else ""
            if not sohd:
                continue
            rows.append({"imei": imei, "phone": phone, "sohd": sohd,
                         "gia_khai": gia, "model": model})
    return rows


# -------------------------------------------------------------------- lookup
def parse_items(url):
    """Tải XML hóa đơn -> list mặt hàng {ten, thtien (chưa VAT), rate, imei}."""
    try:
        r = requests.get(url, timeout=25)
        r.encoding = "utf-8"
        root = ET.fromstring(r.text)
    except Exception:
        return []
    out = []
    for hh in root.iter("HHDVu"):
        ten = hh.findtext("THHDVu") or ""
        thtien = float(hh.findtext("ThTien") or 0)
        mt = re.search(r"(\d+)", hh.findtext("TSuat") or "")
        rate = float(mt.group(1)) / 100.0 if mt else 0.0
        g = ""
        for tin in hh.iter("TTin"):            # IMEI nằm trong ghi chú GChu
            if (tin.findtext("TTruong") or "") == "GChu":
                g = tin.findtext("DLieu") or ""
        mm = re.search(r"\(([A-Za-z0-9]{6,})\)", g)
        out.append({"ten": ten, "thtien": thtien, "rate": rate,
                    "imei": mm.group(1) if mm else ""})
    return out


def lookup(phone, sohd, stop, max_try=40):
    """Tra 1 hóa đơn. Trả (status, [url xml]). status: OK/NOTFOUND/NO_XML/CAPTCHA_FAIL/STOP."""
    bad = 0
    for _ in range(max_try):
        if stop.is_set():
            return "STOP", []
        s = requests.Session()
        s.headers.update({"User-Agent": UA, "Referer": HOME})
        try:
            img = s.get(HOME + f"home/getcaptchaimage?prefix={random.randint(1, 9999999999)}",
                        timeout=20).content
            cap = solve(img)
            bad = bad + 1 if len(cap) != 6 else 0
            # captcha chuẩn 6 ký tự; đọc hụt 4 lần liên tiếp thì nới ra 5-7 để khỏi kẹt
            if len(cap) != 6 and not (bad >= 4 and 5 <= len(cap) <= 7):
                continue
            t = s.post(HOME + "Home/ListInvoice", timeout=20,
                       data={"phone": phone, "invoiceNum": sohd, "captcha": cap,
                             "isCallFromAuto": ""},
                       headers={"X-Requested-With": "XMLHttpRequest"}).text
            if "Mã xác nhận không đúng" in t:
                continue
            urls = re.findall(r'href="(https://einvoices[^"]+)"[^>]*>\s*TẢI XML', t)
            if urls:
                return "OK", list(dict.fromkeys(urls))
            if "không tìm thấy" in t or "không có hóa đơn" in t:
                return "NOTFOUND", []
            return "NO_XML", []
        except Exception:
            continue
    return "CAPTCHA_FAIL", []


def process_group(args):
    """Một (SĐT, số HĐ) -> tra 1 lần, rồi map từng máy trong nhóm vào mặt hàng của hóa đơn."""
    (phone, sohd), devs, stop = args
    status, urls = lookup(phone, sohd, stop)
    pool = []
    for u in urls:
        pool.extend(parse_items(u))
    res = []
    for d in devs:
        fimei = norm(d["imei"])
        match, how = None, ""
        cands = [it for it in pool if it["imei"] and norm(it["imei"]) == fimei]
        if cands:
            match = max(cands, key=lambda x: x["thtien"])
            how = "IMEI"
        else:
            paid = [it for it in pool if it["thtien"] > 0]
            if paid:
                match = max(paid, key=lambda x: x["thtien"])
                how = "Giá cao nhất (IMEI không khớp – rà tay)"
        if match:
            gom = round(match["thtien"] * (1 + match["rate"]))
            gk = float(re.sub("[^0-9]", "", str(d["gia_khai"])) or 0)
            diff = gk - gom
            res.append({**d, "status": status, "cach_match": how,
                        "imei_xml": match["imei"], "ten_xml": match["ten"],
                        "gia_chuaVAT": int(match["thtien"]), "gia_gomVAT": gom,
                        "chenh_lech": int(diff),
                        "ket_qua": "KHỚP" if abs(diff) <= 2 else "LỆCH"})
        else:
            res.append({**d, "status": status, "cach_match": "", "imei_xml": "",
                        "ten_xml": "", "gia_chuaVAT": "", "gia_gomVAT": "",
                        "chenh_lech": "", "ket_qua": "KHÔNG TRA ĐƯỢC"})
    return res


def run_lookup(rows, threads, stop, progress_cb):
    """Gom theo (SĐT, số HĐ) để mỗi hóa đơn chỉ tra 1 lần, chạy song song `threads` luồng."""
    groups = {}
    for r in rows:
        groups.setdefault((r["phone"], r["sohd"]), []).append(r)
    items = [(k, v, stop) for k, v in groups.items()]
    total, done, allres = len(rows), 0, []
    with ThreadPoolExecutor(max_workers=threads) as ex:
        futs = [ex.submit(process_group, it) for it in items]
        for fu in futs:
            if stop.is_set():
                break
            part = fu.result()
            allres.extend(part)
            done += len(part)
            progress_cb(done, total, part)
    return allres


# -------------------------------------------------------------------- output
HEADERS = ["STT", "IMEI/Serial (file)", "SĐT", "Số HĐ", "Model (file)", "Giá KH khai",
           "IMEI trên HĐ", "Tên hàng trên HĐ", "Giá HĐ (chưa VAT)", "Giá HĐ (gồm VAT)",
           "Chênh lệch", "Kết quả"]
KEYS = ["stt", "imei", "phone", "sohd", "model", "gia_khai", "imei_xml", "ten_xml",
        "gia_chuaVAT", "gia_gomVAT", "chenh_lech", "ket_qua"]


def write_excel(path, results):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Đối soát HĐ"
    slate = PatternFill("solid", fgColor="475569")
    white = Font(color="FFFFFF", bold=True)
    ws.merge_cells("A1:L1")
    ws["A1"] = "ĐỐI SOÁT GIÁ TRỊ KH MUA - HÓA ĐƠN ĐIỆN TỬ TGDĐ/ĐMX"
    ws["A1"].fill = slate
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    for c, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.fill = slate
        cell.font = white
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fills = {"KHỚP": PatternFill("solid", fgColor="E7F4EA"),
             "LỆCH": PatternFill("solid", fgColor="FDECEA"),
             "KHÔNG TRA ĐƯỢC": PatternFill("solid", fgColor="FFF4D6")}
    fonts = {"KHỚP": Font(color="1A8245", bold=True),
             "LỆCH": Font(color="C32B3B", bold=True),
             "KHÔNG TRA ĐƯỢC": Font(color="B25E00", bold=True)}
    thin = Side(style="thin", color="E2E6EC")
    for i, r in enumerate(results, start=1):
        rr = i + 2
        r = {**r, "stt": i}
        for c, k in enumerate(KEYS, start=1):
            cell = ws.cell(row=rr, column=c, value=r.get(k, ""))
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if k in ("gia_khai", "gia_chuaVAT", "gia_gomVAT", "chenh_lech"):
                try:
                    cell.value = int(re.sub("[^0-9-]", "", str(r.get(k) or "")))
                    cell.number_format = "#,##0"
                except Exception:
                    pass
        kq = r.get("ket_qua", "")
        for c in range(1, 13):
            if kq in fills:
                ws.cell(row=rr, column=c).fill = fills[kq]
        ws.cell(row=rr, column=12).font = fonts.get(kq, Font())
    for c, w in enumerate([5, 26, 12, 11, 30, 13, 22, 34, 15, 15, 13, 16], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = "A3"
    wb.save(path)
