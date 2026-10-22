# -*- coding: utf-8 -*-
"""
Salvage reconciler engine — faithful Python port of the VBA modules `ThanhLy`
(Celebi reconciler) + the phone-classification bits of `PhoneSalvage`, from
PVI-CELEBI-ECOSWAP_v5.xlsm.

Design: the RATE CARDS stay in the .xlsm (Nam maintains them in Excel). This engine
READS them at runtime (RATE_TBL, LAPTOP_TBL, _Notes, Rate Phone (Celebi), Phone Map)
and only reimplements the LOGIC. So a rate-card edit in Excel is picked up with no
code change, and numbers never diverge from the workbook.

Verified against the workbook's own computed columns (see parity_test.py).
"""
import math, unicodedata, datetime
import openpyxl

CHK = chr(10003)   # ✓
XMK = chr(10007)   # ✗

# ----------------------------------------------------------------- text utils
def strip_accents(s):
    if s is None:
        return ""
    s = str(s).lower()
    out = []
    for ch in unicodedata.normalize("NFD", s):
        if unicodedata.category(ch) == "Mn":     # combining marks
            continue
        out.append(ch)
    return "".join(out).replace("đ", "d").replace("Đ".lower(), "d")

def norm_imei(v):
    s = "" if v is None else str(v)
    p = s.find("\n")
    if p >= 0: s = s[:p]
    p = s.find("(")
    if p >= 0: s = s[:p]
    return s.strip().replace(" ", "").upper()

def imei_text(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    if isinstance(v, int): return str(v)
    return str(v)

def _num(v):
    if v is None: return None
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace(",", "").strip())
    except Exception: return None

def nz(v):
    n = _num(v)
    return n if n is not None else 0.0

def is_num(v):
    return _num(v) is not None and not isinstance(v, str) or (isinstance(v, str) and _num(v) is not None)

def _isnum(v):
    return _num(v) is not None

def parse_vn_date(v):
    if v is None or v == "": return None
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    s = str(v).strip()
    if not s: return None
    if "/" in s:
        p = s.split("/")
        if len(p) == 3:
            try: return datetime.date(int(p[2]), int(p[1]), int(p[0]))
            except Exception: return None
    if "-" in s:
        p = s[:10].split("-")
        if len(p) == 3:
            try: return datetime.date(int(p[0]), int(p[1]), int(p[2]))
            except Exception: return None
    return None

# ----------------------------------------------------------------- classifiers
def usage_idx(m):
    if m < 1: return 0
    if m < 3: return 1
    if m < 6: return 2
    if m < 9: return 3
    return 4

def lap_tier(mn):
    if "macbook" in mn: return 1
    if "dell" in mn or " hp" in mn or "hewlett" in mn: return 2
    if "lenovo" in mn or "acer" in mn or "asus" in mn: return 3
    return 4

def parse_loai(s):
    p = s.find("loai")
    if p >= 0:
        for i in range(p + 4, min(p + 8, len(s) - 1) + 1):
            ch = s[i] if i < len(s) else ""
            if "1" <= ch <= "3": return int(ch)
    return ""

def nhom_dev_type(s):
    t = strip_accents(s)
    if "may tinh bang" in t or "tablet" in t: return "May tinh bang"
    if "laptop" in t: return "Laptop"
    if "dong ho" in t or "smartwatch" in t: return "Dong ho thong minh"
    return ""

def is_phone_model(mn):
    for bad in ("ipad", "tablet", "may tinh bang", "tab "):
        if bad in mn: return False
    for bad in ("laptop", "macbook", "notebook"):
        if bad in mn: return False
    for bad in ("watch", "dong ho", "buds", "airpod"):
        if bad in mn: return False
    for b in ("iphone","galaxy","samsung","oppo","xiaomi","redmi","poco","realme",
              "vivo","nokia","infinix","tecno","honor","huawei","rog phone"):
        if b in mn: return True
    return False

# --------------------------------------------------- phone group (with VIVO fix)
def _cap_of(t):
    if "1tb" in t or "1 tb" in t: return "1TB"
    if "2tb" in t: return "1TB"
    if "512" in t: return "512GB"
    if "256" in t: return "256GB"
    if "128" in t: return "128GB"
    return ""

def _iphone_gen(t):
    p = t.find("iphone")
    if p < 0: return 0
    i = p + 6
    while i < len(t) and t[i] == " ": i += 1
    num = ""
    while i < len(t) and t[i].isdigit():
        num += t[i]; i += 1
    return int(num) if num else 0

def _has_samsung_s(t):
    for i in range(22, 31):
        if f" s{i} " in t: return True
    return False

def dt_nhom(model, gia, nguong=15000000.0):
    t = " " + strip_accents(model) + " "
    if "iphone" in t:
        if " se " in t or "mini" in t: return "iPhone SE/ Mini"
        cap = _cap_of(t); gen = _iphone_gen(t)
        if gen >= 14 and cap: return "iPhone 14, 15 " + cap
        return "iPhone khác"
    if "samsung" in t or "galaxy" in t:
        if "galaxy z" in t or "fold" in t or " flip" in t: return "Samsung Z series"
        if "galaxy s" in t: return "Samsung S series"
        if "galaxy a" in t or "galaxy m" in t: return "Samsung A,M series"
        return "Samsung A,M series"
    if _has_samsung_s(t): return "Samsung S series"
    if "oppo" in t:
        return "Oppo flagship" if gia > nguong else "Oppo, Xiaomi khác"
    if "xiaomi" in t or "redmi" in t or "poco" in t or "vivo" in t:   # VIVO fix
        return "Xiaomi flagship" if gia > nguong else "Oppo, Xiaomi khác"
    return "Android khác"

def celebi_row_for_group(grp, t):
    if grp == "Samsung S series": return "Samsung S series"
    if grp == "Samsung A,M series": return "Samsung A, M series"
    if grp == "Samsung Z series": return "Samsung Z series"
    if grp == "Oppo flagship": return "Oppo > 15tr (Oppo Flagship)"
    if grp == "Xiaomi flagship": return "Xiaomi > 15tr (Xiaomi Flagship)"
    if grp == "iPhone khác": return "iPhone khác"
    if grp == "Oppo, Xiaomi khác":
        if "oppo" in t: return "Oppo khác"
        if "xiaomi" in t or "redmi" in t or "poco" in t or "vivo" in t: return "Xiaomi khác"  # VIVO fix
        return "Android khác"
    return "Android khác"

# ----------------------------------------------------------------- rate cards
class Cards:
    """Rate cards + notes loaded once from the .xlsm."""
    def __init__(self, rate_tbl, laptop_tbl, notes, phone_cel, nguong,
                 phone_eco=None, vat=0.08, loai=None):
        self.rate = rate_tbl        # list of rows; each row is list[17] (cols A..Q, 0-based)
        self.lap = laptop_tbl       # list of 4 rows; each list of month rates (1..N)
        self.notes = notes          # dict code->text
        self.phone = phone_cel      # Celebi phone card: list of (name, [14 rates])
        self.nguong = nguong
        self.phone_eco = phone_eco or []   # Ecoswap phone card: list of (name, [14 rates], house)
        self.vat = vat                     # phone VAT (Phone Map)
        self.loai = loai or {1: 1.0, 2: 0.9, 3: 0.75}   # %Loại factors (Phone Map)

def load_cards(xlsm_path):
    wb = openpyxl.load_workbook(xlsm_path, data_only=True, read_only=True)
    def cell(ws, r, c):
        v = ws.cell(row=r, column=c).value
        return "" if v is None else v
    # RATE_TBL = Rate Card!A5:Q44
    rc = wb["Rate Card"]
    rate = []
    for r in range(5, 45):
        row = [rc.cell(row=r, column=c).value for c in range(1, 18)]  # A..Q
        rate.append(row)
    # LAPTOP_TBL = Laptop Rates!B3:P6  (4 tiers x 15 months)
    lr = wb["Laptop Rates"]
    lap = []
    for r in range(3, 7):
        lap.append([lr.cell(row=r, column=c).value for c in range(2, 17)])  # B..P
    # _Notes
    nws = wb["_Notes"]; notes = {}
    for r in range(2, nws.max_row + 1):
        k = nws.cell(row=r, column=1).value
        if k not in (None, ""):
            notes[str(k)] = "" if nws.cell(row=r, column=2).value is None else str(nws.cell(row=r, column=2).value)
    # Rate Phone (Celebi): col5 name rows 4.., cols 6..19 = T1..T14
    pc = wb["Rate Phone (Celebi)"]
    last = pc.max_row
    phone = []
    for r in range(4, last + 1):
        nm = pc.cell(row=r, column=5).value
        if nm in (None, ""): continue
        rates = [pc.cell(row=r, column=6 + j).value for j in range(14)]
        phone.append((str(nm), rates))
    # Rate Phone (Ecoswap): col5 name rows 4.., cols 6..19 = T1..T14, col20 = Nhà HĐ
    phone_eco = []
    try:
        pe = wb["Rate Phone (Ecoswap)"]
        for r in range(4, pe.max_row + 1):
            nm = pe.cell(row=r, column=5).value
            if nm in (None, ""): continue
            rates = [pe.cell(row=r, column=6 + j).value for j in range(14)]
            house = pe.cell(row=r, column=20).value
            phone_eco.append((str(nm), rates, "" if house is None else str(house).strip()))
    except Exception:
        pass
    # Phone Map: NGUONG_FLAGSHIP, VAT, %LOAI_1/2/3
    nguong, vat, loai = 15000000.0, 0.08, {1: 1.0, 2: 0.9, 3: 0.75}
    try:
        pm = wb["Phone Map"]
        for r in range(4, 41):
            k = str(pm.cell(row=r, column=1).value); v = _num(pm.cell(row=r, column=2).value)
            if v is None: continue
            if k == "NGUONG_FLAGSHIP": nguong = v
            elif k == "VAT": vat = v
            elif k == "%LOAI_1": loai[1] = v
            elif k == "%LOAI_2": loai[2] = v
            elif k == "%LOAI_3": loai[3] = v
    except Exception:
        pass
    wb.close()
    return Cards(rate, lap, notes, phone, nguong, phone_eco=phone_eco, vat=vat, loai=loai)

def dt_rate_celebi(cards, model, months, gia=0.0):
    if not _isnum(months): return None
    mc = int(months)
    if mc < 1: mc = 1
    if mc > 14: mc = 14
    t = " " + strip_accents(model) + " "
    best_i, best_len = -1, 0
    for i, (nm, rates) in enumerate(cards.phone):
        cm = strip_accents(nm)
        if cm and cm in t and len(cm) > best_len:
            best_len = len(cm); best_i = i
    if best_i >= 0:
        return cards.phone[best_i][1][mc - 1]
    grp = dt_nhom(model, gia, cards.nguong)
    cel = strip_accents(celebi_row_for_group(grp, t))
    for nm, rates in cards.phone:
        if strip_accents(nm) == cel:
            return rates[mc - 1]
    return None

# ----------------------------------------------------------------- helpers on rate row
def _rt_use(row, ui):   # cols J..N = index 9..13
    v = row[9 + ui]
    return v if _isnum(v) else 0
def _rt_loai(row, ln):  # cols O..Q = index 14..16
    v = row[14 + ln - 1]
    return v if _isnum(v) else 1
def _tier_name(row):
    b = strip_accents(row[3])
    if not b.strip(): return "Toan bo"
    if "khac" in b: return "Hang khac"
    return "Hang nhom 1"
def _root_cause(rate, chosen_row, ui, bb_use):
    row = rate[chosen_row]
    if ui - 1 >= 0 and abs(nz(_rt_use(row, ui - 1)) - bb_use) < 1e-4: return "BOUND"
    if ui + 1 <= 4 and abs(nz(_rt_use(row, ui + 1)) - bb_use) < 1e-4: return "BOUND"
    d = row[1]
    for j in range(len(rate)):
        if rate[j][1] == d:
            for k in range(5):
                v = rate[j][9 + k]
                if _isnum(v) and abs(float(v) - bb_use) < 1e-4: return "BRAND"
    return "RATEERR"

def make_note(nc, notes):
    if nc.startswith("MISS:"):
        q = '"'
        parts = [p for p in nc[5:].split(",") if p]
        lst = ", ".join(q + str(notes.get(p, p)) + q for p in parts)
        return f"{notes.get('MISS_PRE','')} {lst} {notes.get('MISS_SUF','')}"
    return notes.get(nc, nc)

# ----------------------------------------------------------------- ComputeRow
def compute_row(model, imei_raw, def_nhom, nhom_hang, cards, ref):
    """ref = dict for this IMEI: price, mua, gs, namr, loai, bbVal, bbUse (or {} if absent)."""
    rate = cards.rate
    imei = norm_imei(imei_raw)
    mn = " " + strip_accents(model) + " "
    # 1) keyword hit (longest keyword across all rows)
    hit, best_len = 0, 0
    for i in range(len(rate)):
        for kw in str(rate[i][2]).split(", "):
            if kw and kw in mn and len(kw) > best_len:
                best_len = len(kw); hit = i + 1     # 1-based like VBA
    # 2) device type
    if def_nhom == 3:
        dt = nhom_dev_type(nhom_hang)
        if not dt:
            dt = "Dien thoai di dong" if is_phone_model(mn) else (str(rate[hit-1][1]) if hit > 0 else "Laptop")
    elif hit == 0:
        return dict(nhomTxt="?", dt="?", tier="", months="", usage="", namUsed="",
                    userate="", loaif="", formula="", auto="", bbVal=ref.get("bbVal",""),
                    diff="", symx=XMK + " Mismatch", code="UNKNOWN", price=ref.get("price",0))
    else:
        dt = str(rate[hit-1][1])
    # 3) choose the brand/tier row within this device type
    FB, chosen = 0, 0
    for jj in range(1, len(rate) + 1):
        row = rate[jj-1]
        if str(row[1]) == dt:
            tier_sa = strip_accents(row[3])
            is_fb = (not tier_sa.strip()) or ("khac" in tier_sa)
            if is_fb:
                FB = jj
            else:
                for bk in str(row[3]).split(", "):
                    if bk.strip() and strip_accents(bk.strip()) in mn:
                        chosen = jj; break
            if chosen > 0: break
    if chosen == 0:
        chosen = FB if FB > 0 else (hit if hit > 0 else 1)
    crow = rate[chosen-1]
    nhom = int(crow[0])

    price = nz(ref.get("price")) if ref.get("price") not in (None, "") else 0.0
    bbVal = ref.get("bbVal", "")
    bbUse = ref.get("bbUse", "")
    loaiN = ref.get("loai", "")
    loaiN = int(loaiN) if str(loaiN).strip().isdigit() else 3
    if loaiN < 1 or loaiN > 3: loaiN = 3
    namr = ref.get("namr", "")

    mua, gs = ref.get("mua"), ref.get("gs")
    has_month = isinstance(mua, datetime.date) and isinstance(gs, datetime.date)
    months = ((gs - mua).days / 30.0) if has_month else 0.0

    userate = ""; loaif = ""; auto = ""; usage = ""; namUsed = ""
    labs = ["<1 thang", "1-3 thang", "3-6 thang", "6-9 thang", "9-12 thang"]
    lapN = ["Macbook", "Dell, HP", "Lenovo/Acer/Asus", "Laptop khac"]
    nhomTxt = f"{nhom}" + (" - Dien may, dien tu" if nhom == 1 else
                           (" - Dien gia dung" if nhom == 2 else " - Dien tu vien thong"))
    tier = _tier_name(crow); ui = 0

    if dt == "Laptop":
        formula = "LAPMONTH"; lt = lap_tier(mn); tier = lapN[lt-1]; loaif = _rt_loai(crow, loaiN)
        if has_month:
            mcol = math.ceil(months); mcol = max(1, min(mcol, len(cards.lap[0])))  # clamp to table width
            userate = cards.lap[lt-1][mcol-1]; usage = f"Thang {mcol}"
            auto = float(userate) * float(loaif) * price
    elif dt == "Dien thoai di dong":
        formula = "PHONECEL"; tier = "Celebi per-model"; loaif = _rt_loai(crow, loaiN)
        if has_month:
            mcol = math.ceil(months); mcol = max(1, min(mcol, 14)); usage = f"Thang {mcol}"
            gr = dt_rate_celebi(cards, model, mcol, price)
            if _isnum(gr):
                userate = float(gr); auto = userate * float(loaif) * price
    else:
        formula = "NAMUSE" if nhom == 1 else ("USE" if nhom == 2 else "USELOAI")
        if has_month:
            ui = usage_idx(months); usage = labs[ui]
            if formula == "USE":
                userate = _rt_use(crow, ui); auto = float(userate) * price
            elif formula == "NAMUSE":
                userate = _rt_use(crow, ui); namUsed = float(namr) if _isnum(namr) else 0
                auto = (float(namUsed) + float(userate)) * price
            else:  # USELOAI
                userate = _rt_use(crow, ui); loaif = _rt_loai(crow, loaiN)
                auto = float(userate) * float(loaif) * price

    # reconcile
    miss = []
    if price == 0 or ref.get("price") in (None, ""):
        if ref.get("price") in (None, "") or str(ref.get("price")).strip() == "": miss.append("COL_PRICE")
    if bbVal in (None, "") or str(bbVal).strip() == "": miss.append("COL_VAL")
    if not (isinstance(mua, datetime.date)): miss.append("COL_MUA")
    if not (isinstance(gs, datetime.date)): miss.append("COL_GS")
    if formula == "NAMUSE" and (namr in (None, "") or str(namr).strip() == ""): miss.append("COL_NAM")

    symx = ""; code = ""; diff = ""
    if miss:
        symx = XMK + " Mismatch"; code = "MISS:" + ",".join(miss)
    elif not _isnum(auto) or not _isnum(bbVal):
        symx = "-"; code = "NODATA"
    else:
        diff = float(auto) - float(bbVal); thr = max(0.005 * abs(float(bbVal)), 1.0)
        if abs(diff) <= thr:
            symx = CHK + " Match"
        else:
            symx = XMK + " Mismatch"
            has_exp = False; expv = 0.0
            if _isnum(bbUse):
                expv = ((nz(namr) + float(bbUse)) * price) if nhom == 1 else (float(bbUse) * price)
                has_exp = True
            if has_exp and abs(float(bbVal) - expv) > thr:
                code = "PASTE"
            elif dt == "Laptop":
                code = "LAPDIFF"
            elif nhom == 3:
                code = "LAPHIGH" if float(auto) < float(bbVal) else "LAPLOW"
            elif _isnum(bbUse):
                code = _root_cause(rate, chosen-1, ui, float(bbUse))
            else:
                code = "RATEERR"

    return dict(nhomTxt=nhomTxt, dt=str(crow[1]), tier=tier,
                months=round(months, 2) if has_month else "", usage=usage, namUsed=namUsed,
                userate=userate, loaif=loaif, formula=formula, auto=auto, bbVal=bbVal,
                diff=diff, symx=symx, code=code, price=price)

# convenience: fixed price-column note label text (formula) via notes
def formula_label(code_or_formula, cards):
    return cards.notes.get(code_or_formula, code_or_formula)

# ============================================================ ECOSWAP 1-đổi-1
def round_hu(x):
    """Round half away-from-zero (Excel ROUND) with tiny epsilon, like VBA RoundHU."""
    if x >= 0:
        return math.floor(x + 0.500001)
    return -math.floor(-x + 0.500001)

def dt_loai(gsc, cards):
    """GSC 'Loại 1/2/3' -> %Loại factor (or None)."""
    s = strip_accents(gsc)
    p = s.find("loai")
    if p >= 0:
        for i in range(p + 4, min(p + 8, len(s) - 1) + 1):
            ch = s[i] if i < len(s) else ""
            if ch in ("1", "2", "3"):
                return cards.loai.get(int(ch))
    return None

def dt_thang(mua, gs):
    """Months used (ceiling, min 1) from purchase->GS dates; None if invalid."""
    if not (isinstance(mua, datetime.date) and isinstance(gs, datetime.date)):
        return None
    d = (gs - mua).days / 30.0
    if d < 0: d = 0.0
    m = math.ceil(d)
    return m if m >= 1 else 1

def _phone_idx(names, model, nhom):
    """Longest col-name substring in ' model ', else exact match on CelebiRowForGroup(nhom)."""
    t = " " + strip_accents(model) + " "
    best_i, best_len = -1, 0
    for i, nm in enumerate(names):
        cm = strip_accents(nm)
        if cm and cm in t and len(cm) > best_len:
            best_len = len(cm); best_i = i
    if best_i >= 0:
        return best_i
    cel = strip_accents(celebi_row_for_group(nhom, t))
    for i, nm in enumerate(names):
        if strip_accents(nm) == cel:
            return i
    return -1

def compute_eco_row(model, gia, si, gsc, mua, gs, cards):
    gia = nz(gia); si = nz(si)
    nhom = dt_nhom(model, gia, cards.nguong)
    loaiV = dt_loai(gsc, cards)
    loai = float(loaiV) if loaiV is not None else 0.0
    th = dt_thang(mua, gs)
    has_th = th is not None
    m = min(max(th, 1), 14) if has_th else 0

    grEco = -1.0; house = ""
    if has_th:
        idx = _phone_idx([e[0] for e in cards.phone_eco], model, nhom)
        if idx >= 0:
            v = cards.phone_eco[idx][1][m - 1]
            if _isnum(v): grEco = float(v)
            house = cards.phone_eco[idx][2]
    grCel = -1.0
    if has_th:
        idxc = _phone_idx([c[0] for c in cards.phone], model, nhom)
        if idxc >= 0:
            v = cards.phone[idxc][1][m - 1]
            if _isnum(v): grCel = float(v)

    ecoGross = ecoNet = ecoVat = None
    if grEco >= 0 and loai > 0:
        ecoGross = round_hu(loai * grEco * gia)
        ecoNet = round_hu(ecoGross / (1 + cards.vat))
        ecoVat = ecoGross - ecoNet
    celGross = None
    if grCel >= 0 and loai > 0:
        celGross = round_hu(loai * grCel * gia)
    check = (si * 0.9 - gia) if (si > 0 or gia > 0) else None
    if ecoGross is not None and celGross is not None:
        higher = "Ecoswap" if ecoGross >= celGross else "Celebi"
    elif ecoGross is not None:
        higher = "Ecoswap"
    elif celGross is not None:
        higher = "Celebi"
    else:
        higher = ""
    return dict(nhom=nhom, loai=(loai if loai > 0 else ""), thang=(th if has_th else ""),
                grEco=(grEco if grEco >= 0 else ""), ecoNet=(ecoNet if ecoNet is not None else ""),
                ecoVat=(ecoVat if ecoVat is not None else ""),
                ecoGross=(ecoGross if ecoGross is not None else ""),
                grCel=(grCel if grCel >= 0 else ""), celGross=(celGross if celGross is not None else ""),
                check=(check if check is not None else ""), house=house, higher=higher)

# ---- import a DT-sheet phone list (DoImportEco port) ----
def _find_col_rows(ws, key, r1, r2, maxc=50):
    for c in range(1, maxc + 1):
        for r in range(r1, r2 + 1):
            if key in strip_accents(ws.cell(row=r, column=c).value):
                return c
    return 0

def _find_dt_sheet(wb):
    for ws in wb.worksheets:
        if strip_accents(ws.title) == "dt":
            return ws
    for ws in wb.worksheets:
        if _find_col_rows(ws, "sum insured", 1, 5) > 0 and _find_col_rows(ws, "imei", 1, 5) > 0:
            return ws
    return None

def import_dt(path):
    """Read an Ecoswap DT-sheet phone list -> list of dict(claim, model, imei, imei2,
       mua, gs, si, gia, gsc, dot)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    dt = _find_dt_sheet(wb)
    if dt is None:
        wb.close()
        raise ValueError("Không tìm thấy sheet DT (hoặc sheet có cột Sum Insured + IMEI).")
    hr = 1
    for rr in range(1, 7):
        if _find_col_rows(dt, "imei", rr, rr) > 0:
            hr = rr; break
    # Header co the 1 hoac 2 tang. Dong ngay duoi dong "IMEI" chi la header phu khi o Model/IMEI
    # cua no trong; co model/IMEI tuc la may dau tien (file dot 55-61: header 1 dong, may o dong 3
    # bi bo sot khi con mac dinh 2 dong).
    r1 = r2 = hr
    cM0, cI0 = _find_col_rows(dt, "model", hr, hr), _find_col_rows(dt, "imei", hr, hr)
    if not any(str(dt.cell(row=hr + 1, column=c).value or "").strip() for c in (cM0, cI0) if c):
        r2 = hr + 1
    def fc(*keys):
        for k in keys:
            c = _find_col_rows(dt, k, r1, r2)
            if c: return c
        return 0
    cClaim = fc("claim request"); cModel = fc("model")
    cImei = fc("imei/serial", "imei"); cImei2 = fc("change warranty")
    cMua = fc("ngay mua"); cGS = fc("gs gui", "gui yeu cau thanh ly")
    cSI = fc("sum insured"); cGia = fc("gia tri may kh mua", "kh mua")
    cGsc = fc("gsc phan loai", "phan loai"); cDot = fc("dot")
    def cell(r, c):
        v = dt.cell(row=r, column=c).value
        return "" if v is None else v
    key_col = cModel if cModel else cImei
    last = dt.max_row
    out = []
    for rr in range(r2 + 1, last + 1):
        mdl = str(cell(rr, cModel)) if cModel else ""
        imv = cell(rr, cImei) if cImei else ""
        if str(mdl).strip() == "" and str(imv).strip() == "":
            continue
        out.append(dict(
            claim=cell(rr, cClaim) if cClaim else "",
            model=mdl, imei=imei_text(imv),
            imei2=imei_text(cell(rr, cImei2)) if cImei2 else "",
            mua=parse_vn_date(cell(rr, cMua)) if cMua else None,
            gs=parse_vn_date(cell(rr, cGS)) if cGS else None,
            si=cell(rr, cSI) if cSI else "", gia=cell(rr, cGia) if cGia else "",
            gsc=cell(rr, cGsc) if cGsc else "", dot=cell(rr, cDot) if cDot else "",
        ))
    wb.close()
    return out

def reconcile_eco_file(path, cards):
    devices = import_dt(path)
    out = []
    for i, d in enumerate(devices, 1):
        res = compute_eco_row(d["model"], d["gia"], d["si"], d["gsc"], d["mua"], d["gs"], cards)
        row = dict(d); row["stt"] = i; row.update(res)
        out.append(row)
    return out

ECO_HEADERS = ["STT", "Số đơn (Claim)", "Model", "IMEI/Serial", "IMEI mới", "Ngày mua",
    "Ngày GS gửi", "Sum Insured", "Giá KH mua", "GSC phân loại", "Đợt", "Phân nhóm",
    "%Loại", "Số tháng", "%GRV Eco", "Giá Eco (-VAT)", "VAT", "Giá Eco (+VAT)",
    "%GRV Celebi", "Giá Celebi", "Check giá KH", "Nhà thanh lý theo HĐ", "Nhà thanh lý cao hơn"]

def eco_record(r):
    return [r.get("stt", ""), r.get("claim", ""), r.get("model", ""), r.get("imei", ""),
            r.get("imei2", ""), r.get("mua"), r.get("gs"), r.get("si", ""), r.get("gia", ""),
            r.get("gsc", ""), r.get("dot", ""), r.get("nhom", ""), r.get("loai", ""),
            r.get("thang", ""), r.get("grEco", ""), r.get("ecoNet", ""), r.get("ecoVat", ""),
            r.get("ecoGross", ""), r.get("grCel", ""), r.get("celGross", ""), r.get("check", ""),
            r.get("house", ""), r.get("higher", "")]

def summary_eco(results):
    n = len(results)
    nEco = sum(1 for r in results if str(r.get("house", "")) == "Ecoswap")
    nCel = sum(1 for r in results if str(r.get("house", "")) == "Celebi")
    sEco = sum(float(r["ecoGross"]) for r in results if _isnum(r.get("ecoGross")))
    sCel = sum(float(r["celGross"]) for r in results if _isnum(r.get("celGross")))
    nHigherEco = sum(1 for r in results if r.get("higher") == "Ecoswap")
    nHigherCel = sum(1 for r in results if r.get("higher") == "Celebi")
    return dict(n=n, nEco=nEco, nCel=nCel, sEco=sEco, sCel=sCel,
                nHigherEco=nHigherEco, nHigherCel=nHigherCel)

def build_eco_workbook(results, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "Ecoswap"
    hf = PatternFill("solid", fgColor="1F4E79"); hfont = Font(bold=True, color="FFFFFF", name="Calibri")
    for j, h in enumerate(ECO_HEADERS, 1):
        c = ws.cell(row=1, column=j, value=h); c.fill = hf; c.font = hfont
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, r in enumerate(results, start=2):
        for j, v in enumerate(eco_record(r), 1):
            cell = ws.cell(row=i, column=j, value=(v if v != "" else None))
            if j in (4, 5):                       # IMEI / IMEI mới -> force TEXT (no 8.6e+11)
                cell.number_format = "@"
                if v not in ("", None): cell.value = str(v)
        for col in (8, 9, 16, 17, 18, 20, 21): ws.cell(row=i, column=col).number_format = "#,##0"
        for col in (6, 7): ws.cell(row=i, column=col).number_format = "dd/mm/yyyy"
        ws.cell(row=i, column=13).number_format = "0%"
        ws.cell(row=i, column=15).number_format = "0.0%"
        ws.cell(row=i, column=19).number_format = "0.0%"
    widths = [5, 16, 30, 16, 14, 11, 11, 12, 12, 12, 6, 18, 7, 7, 9, 13, 10, 13, 10, 13, 13, 14, 14]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    wb.save(out_path)
    return out_path

# ----------------------------------------------------------------- BBDS import
def _find_col(ws, key, hr, maxc=90):
    for c in range(1, maxc + 1):
        if key in strip_accents(ws.cell(row=hr, column=c).value):
            return c
    return 0

def _hdr_model_col(ws, hr):
    c = _find_col(ws, "kieu thiet bi", hr)
    if c == 0: c = _find_col(ws, "ten san pham", hr)
    return c

def _find_bbds_sheet(wb):
    for ws in wb.worksheets:
        if "bbds" in strip_accents(ws.title):
            return ws
    for ws in wb.worksheets:
        for r in range(1, 61):
            a = strip_accents(ws.cell(row=r, column=1).value)
            if "nhom hang" in a:
                return ws
            if _hdr_model_col(ws, r) > 0 and _find_col(ws, "gia tri thu cu", r) > 0:
                return ws
    return None

def import_bbds(path):
    """Port of DoImport: read a BBDS workbook -> (devices, ref_by_imei).
       devices: list of dict(stt, model, imei, nhomHang, tinhtrang, ttdot).
       ref_by_imei: dict normimei -> dict(sec, price, mua, gs, namr, loai, bbVal, bbUse).
       Note: openpyxl reads ALL rows regardless of AutoFilter, so the .End(xlUp)
       filter-hidden-rows trap that bit the VBA cannot happen here."""
    wb = openpyxl.load_workbook(path, data_only=True)
    bbs = _find_bbds_sheet(wb)
    if bbs is None:
        wb.close()
        raise ValueError("Không tìm thấy sheet BBDS trong file đã chọn.")
    def cell(r, c):
        v = bbs.cell(row=r, column=c).value
        return "" if v is None else v
    sec = 0
    cStt = cKieu = cImei = cNhom = cMua = cGS = cGia = cTinh = cNam = cUse = cVal = cDot = cLoai = 0
    devices, ref = [], {}
    for rr in range(1, bbs.max_row + 1):
        a1 = strip_accents(cell(rr, 1))
        if "nhom hang dien gia dung" in a1:
            sec = 1
        elif "nhom hang dien may" in a1:
            sec = 2
        elif "nhom hang dien thoai" in a1:
            sec = 3
        elif "nhom hang dien tu" in a1:
            sec = 3
        elif _hdr_model_col(bbs, rr) > 0:
            cStt = _find_col(bbs, "stt", rr); cKieu = _hdr_model_col(bbs, rr); cImei = _find_col(bbs, "imei", rr)
            cLoai = _find_col(bbs, "loai thiet bi", rr)
            cNhom = _find_col(bbs, "nhom hang", rr); cMua = _find_col(bbs, "ngay khach hang mua", rr)
            cGS = _find_col(bbs, "gs thong bao", rr); cGia = _find_col(bbs, "gia khach hang mua", rr)
            cTinh = _find_col(bbs, "tinh trang", rr); cNam = _find_col(bbs, "ty le nam", rr)
            cUse = _find_col(bbs, "ty le thoi gian", rr); cVal = _find_col(bbs, "gia tri thu cu", rr)
            cDot = _find_col(bbs, "tt dot", rr)
        elif cImei > 0:
            imv = cell(rr, cImei)
            sttv = cell(rr, cStt if cStt > 0 else 1)
            if str(imv).strip() != "" and _isnum(sttv):
                model = cell(rr, cKieu)
                key = norm_imei(imei_text(imv))
                loai = parse_loai(strip_accents(cell(rr, cLoai))) if cLoai > 0 else parse_loai(strip_accents(cell(rr, cNhom)))
                devices.append(dict(stt=sttv, model=str(model), imei=imei_text(imv),
                                    nhomHang=cell(rr, cNhom) if cNhom > 0 else "",
                                    tinhtrang=cell(rr, cTinh) if cTinh > 0 else "",
                                    ttdot=cell(rr, cDot) if cDot > 0 else ""))
                ref[key] = dict(sec=sec, price=cell(rr, cGia) if cGia > 0 else "",
                                mua=parse_vn_date(cell(rr, cMua)) if cMua > 0 else None,
                                gs=parse_vn_date(cell(rr, cGS)) if cGS > 0 else None,
                                namr=cell(rr, cNam) if cNam > 0 else "",
                                loai=loai, bbVal=cell(rr, cVal) if cVal > 0 else "",
                                bbUse=cell(rr, cUse) if cUse > 0 else "")
    wb.close()
    return devices, ref

def reconcile_file(bbds_path, cards):
    """End-to-end: import a BBDS file + compute every row. Returns list of result dicts
       (device fields + compute_row outputs)."""
    devices, ref = import_bbds(bbds_path)
    out = []
    for d in devices:
        key = norm_imei(d["imei"])
        rf = ref.get(key, {})
        defN = 3 if (rf.get("sec") == 3 or str(rf.get("sec")) == "3") else 0
        res = compute_row(d["model"], d["imei"], defN, d["nhomHang"], cards, rf)
        row = dict(d); row.update(res)
        out.append(row)
    return out

# ----------------------------------------------------------------- output / summary
OUT_HEADERS = ["STT", "Model", "IMEI/Serial", "Giá mua (BBDS)", "Nhóm", "Loại thiết bị",
    "Bậc hãng", "Số tháng SD", "Bậc thời gian", "Tỷ lệ năm", "Tỷ lệ TG", "Hệ số Loại",
    "Công thức", "GT thu cũ (AUTO)", "GT thu cũ (BBDS)", "Chênh lệch", "Kết quả", "Ghi chú",
    "Nhóm Hàng (BBDS)", "Tình trạng", "TT đợt"]

def result_record(r, cards):
    """One result dict -> ordered values matching OUT_HEADERS."""
    return [
        r.get("stt", ""), r.get("model", ""), r.get("imei", ""), r.get("price", ""),
        r.get("nhomTxt", ""), r.get("dt", ""), r.get("tier", ""), r.get("months", ""),
        r.get("usage", ""), r.get("namUsed", ""), r.get("userate", ""), r.get("loaif", ""),
        cards.notes.get(r.get("formula", ""), r.get("formula", "")),
        r.get("auto", ""), r.get("bbVal", ""), r.get("diff", ""), r.get("symx", ""),
        make_note(str(r.get("code", "")), cards.notes),
        r.get("nhomHang", ""), r.get("tinhtrang", ""), r.get("ttdot", ""),
    ]

def summary(results):
    """KPI by nhóm (1=điện máy/điện tử, 2=điện gia dụng, 3=điện tử viễn thông) + tổng."""
    groups = {1: "Điện máy / Điện tử", 2: "Điện gia dụng", 3: "Điện tử viễn thông"}
    agg = {g: dict(n=0, ok=0, bad=0, auto=0.0, bb=0.0) for g in groups}
    for r in results:
        nh = str(r.get("nhomTxt", "")).strip()
        g = int(nh[0]) if nh[:1] in ("1", "2", "3") else None
        if g is None: continue
        a = agg[g]; a["n"] += 1
        if CHK in str(r.get("symx", "")): a["ok"] += 1
        if XMK in str(r.get("symx", "")): a["bad"] += 1
        if _isnum(r.get("auto")): a["auto"] += float(r["auto"])
        if _isnum(r.get("bbVal")): a["bb"] += float(r["bbVal"])
    rows = []
    tot = dict(n=0, ok=0, bad=0, auto=0.0, bb=0.0)
    for g, label in groups.items():
        a = agg[g]
        for k in tot: tot[k] += a[k]
        rows.append(dict(nhom=label, **a, diff=a["auto"] - a["bb"]))
    rows.append(dict(nhom="TỔNG", **tot, diff=tot["auto"] - tot["bb"]))
    return rows

def build_output_workbook(results, out_path, cards):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "Kết quả"
    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(bold=True, color="FFFFFF", name="Calibri")
    for j, h in enumerate(OUT_HEADERS, 1):
        c = ws.cell(row=1, column=j, value=h); c.fill = hdr_fill; c.font = hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    money = '#,##0'
    for i, r in enumerate(results, start=2):
        vals = result_record(r, cards)
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=j, value=(v if v != "" else None))
            if j == 3 and v not in ("", None):    # IMEI/Serial -> force TEXT
                cell.number_format = "@"; cell.value = str(v)
        for col in (4, 14, 15, 16): ws.cell(row=i, column=col).number_format = money
        ws.cell(row=i, column=8).number_format = "0.00"
        ws.cell(row=i, column=10).number_format = "0%"
        ws.cell(row=i, column=11).number_format = "0.0%"
        ws.cell(row=i, column=12).number_format = "0%"
        sym = str(r.get("symx", "")); sc = ws.cell(row=i, column=17); sc.font = Font(bold=True,
            color=("16A34A" if CHK in sym else ("DC2626" if XMK in sym else "000000")))
    widths = [5, 34, 18, 13, 20, 20, 16, 9, 11, 8, 8, 8, 18, 14, 14, 12, 12, 40, 20, 12, 8]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out_path)
    return out_path
