# -*- coding: utf-8 -*-
"""Search GCNBH (Preview): tu so hop dong trong file claim, tra ra link file giay chung nhan
   tren OneDrive cua anh Thanh (P.TBHH). Xuat FILE RIENG, khong dung toi file _MAPPED.

BAN 30/09/2026 - QUET SONG, KHONG CON INDEX:
   Kho da doi cho sang `PVI/GCN + XMP - DMX/GCNBH + XMP - TGDĐ` va user da dat ten chuan
   `GCN 25 P000123_E01.pdf`. May host sync kho qua OneDrive (Files On-Demand: chi co TEN, 0 byte),
   quet ten ~5.000 file mat ~0,3 s nen moi lan tra la quet lai -> file moi up len la thay ngay
   (tre bang do tre sync cua OneDrive). Cache 30 s chi de do quet lai moi lan bam nut.

NGUYEN TAC:
   - Khoa tra cuu = (nam, chu P/C, so thu tu, so sua doi E). So thu tu LAP LAI moi nam:
     24/45/09/KHAC/P000028 != 25/45/09/KHAC/P000028.
   - MOT so don co the co NHIEU GCN theo san pham (P000158 co ca giay 1D1 lan BHMR, cung STBH/phi).
     San pham lay tu TEN FOLDER. Biet san pham cua ho so thi chi tra giay san pham do.
   - KHONG BAO GIO tu thay bang ban E khac hay san pham khac (Nam chot 28/07/2026):
     in nham giay vao ho so boi thuong con te hon khong tim ra.
   - Ten file chua chuan (032.01, P000857, SDBS 297.01...) van doc theo luat cu nhung ghi chu ro.
"""
import ctypes
import datetime as dt
import glob
import io
import os
import re
import time
import unicodedata
import urllib.parse
from concurrent.futures import ThreadPoolExecutor

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

KHO_TEN = "GCNBH + XMP - TGDĐ"
WEB = ("https://pvicomvn-my.sharepoint.com/personal/thanhnnd_pvi_com_vn/Documents/"
       "PVI/GCN + XMP - DMX/GCNBH + XMP - TGDĐ")
TTL = 30

RE_POL = re.compile(r"(\d{2})/(\d{2})/(\d{2})/KHAC/([PC])(\d{6,7})(?:/E(\d{1,3}))?", re.I)
RE_GSC = re.compile(r"GSC-\d{6}-\d+")

# Ma san pham doc tu ten folder (so tren chuoi da bo dau, bo moi ky tu khong phai chu/so).
# Thu tu quan trong: NVMH truoc (BHNVMH), BHRV truoc (BHRVMH, RVMH).
SP_FOLDER = [
    ("NVMH", ("NVMH", "BHNV", "NUTVOMANHINH")),
    ("BHRV", ("BHRV", "RVMH", "BMRV", "ROIVO")),
    ("1D1",  ("1D1", "1DOI1", "MOTDOIMOT")),
    ("BHMR", ("BHMR", "MORONG")),
    ("MDM",  ("MDM",)),
    ("ADLD", ("ADLD",)),
    ("BHNC", ("BHNC",)),
]
# Ten san pham trong file claim (cot Product, vd "BẢO HÀNH MỞ RỘNG TIVI 12 THÁNG (R)").
SP_CLAIM = [
    ("1D1",  ("1 DOI 1", "MOT DOI MOT", "1D1")),
    ("NVMH", ("NUT VO",)),
    ("BHRV", ("ROI VO",)),
    ("BHMR", ("MO RONG",)),
]
# Goi y theo TEN FILE claim (vd "Đợt 43_184 HS_1 doi 1_T4.xlsx", "DOT 11_140 HS_EW_T02.xlsx").
SP_TENFILE = [
    ("1D1",  ("1 DOI 1", "1DOI1", "1D1")),
    ("NVMH", ("NVMH", "NUT VO")),
    ("BHRV", ("BHRV", "ROI VO")),
    ("BHMR", ("BHMR", "EW", "MO RONG")),
]
SP_LIST = ["1D1", "BHMR", "BHRV", "NVMH", "MDM", "ADLD", "BHNC"]
TU_NHAN, TAT_CA = "Tự nhận theo file", "Tất cả sản phẩm"

# Tu duoc phep dung truoc so thu tu trong ten file kieu cu ("gcnbh 1d1 - 172", "SĐBS 240_E02").
TIEN_TO_OK = {"GCNBH", "GCN", "SDBS", "BH", "TGDD", "1D1", "BHMR", "BHRV", "NVMH", "MDM",
              "ADLD", "BHNC", "BHNV", "RVMH", "GCNBHMR"}


def fold(s):
    s = unicodedata.normalize("NFD", str(s)).replace("đ", "d").replace("Đ", "D")
    return "".join(c for c in s if unicodedata.category(c) != "Mn").upper()


def tight(s):
    return re.sub(r"[^0-9A-Z]", "", fold(s))


def sp_cua_folder(parts):
    """Doc tu folder sau nhat len tren, gap ma san pham dau tien thi lay."""
    for seg in reversed(parts):
        t = tight(seg)
        for code, keys in SP_FOLDER:
            if any(k in t for k in keys):
                return code
    return ""


def doc_ten(stem, yy_folder):
    """Ten file (khong duoi) -> (yy, chu, so, E, chuan?) hoac None neu khong doc duoc."""
    s = re.sub(r"\s*\(\d+\)$", "", stem.strip())
    s = re.sub(r"([PC]\d{3})_(\d{3})\b", r"\1\2", s)     # "P000_329" go nham dau _
    m = re.match(r"^GCN (\d{2}) ([PC])(\d{6})(?:_E(\d{2,3}))?$", s, re.I)
    if m:
        return m.group(1), m.group(2).upper(), int(m.group(3)), int(m.group(4) or 0), True
    s = re.sub(r"^\d{1,3}\.\s+", "", s)        # tien to "12. " (phai co dau cach: "032." la so don)
    m = re.search(r"(\d{2})[._/]\d{2}[._/]\d{2}[._/]KHAC[._/]([PC])(\d{3,7})(?:[._/]E(\d{1,3}))?",
                  s, re.I)
    if m:
        return m.group(1), m.group(2).upper(), int(m.group(3)), int(m.group(4) or 0), False
    m = re.search(r"(?:^|\s)(?:(\d{2})\s+)?([PC])(\d{5,7})(?:[-_\s.]+E?(\d{1,3}))?$", s, re.I)
    if m:
        yy = m.group(1) or yy_folder
        return (yy, m.group(2).upper(), int(m.group(3)), int(m.group(4) or 0), False) if yy else None
    m = re.search(r"(?:^|[\s-])(\d{1,4})(?:[._\s]+E?(\d{1,3}))?$", s, re.I)
    if m and yy_folder:
        truoc = re.sub(r"[\s\-_.]+", " ", fold(s[:m.start(1)])).split()
        if all(w in TIEN_TO_OK or re.fullmatch(r"T\d{1,2}", w) for w in truoc):
            return yy_folder, "P", int(m.group(1)), int(m.group(2) or 0), False
    return None


# ---- XMP (ban xac nhan nop phi): 1 file / thang / san pham / (ban le | chuong trinh tang) --------
# Ben trong la bang ke so don (kem E) nop phi trong thang, nhung la ANH SCAN, khong doc chu duoc.
# Nen ghep theo VI TRI: GCN nam o folder thang nao thi lay XMP cung thang, cung san pham, cung loai
# ban le / tang. Kiem chung: GCN 26 P000155 o 2026/T02.2026/BHMR <-> XMP BHMR T02_2026 co dong P000155.
# MDM va 1D1 dung lan cho nhau (folder "1D1 TANG MAY LANH" chua "XMP BH MDM") nen thieu thi thu cheo.
CHEO = {"1D1": "MDM", "MDM": "1D1"}


def sp_xmp(ten):
    t = tight(ten[3:])                          # bo chu XMP
    for code, keys in SP_FOLDER:
        if any(k in t for k in keys):
            return code
    return ""


def la_tang(rel):
    t = tight("/".join(rel.split("/")[1:-1]) + "/" + rel.split("/")[-1])
    return "TANG" in t or "COMBO" in t


def tim_xmp(xmps, hit):
    """-> [xmp path] ung vien cho 1 file GCN. Di tu folder cua GCN len (dung truoc folder nam),
    folder dau tien co XMP dung san pham thi lay."""
    parts = hit["path"].split("/")[:-1]
    sp = hit["sp"]
    if not sp:
        return []
    tang = la_tang(hit["path"])
    for n in range(len(parts), 1, -1):
        goc = "/".join(parts[:n]) + "/"
        o_day = [x for x in xmps if x["path"].startswith(goc)]
        if not o_day:
            continue
        for thu in (sp, CHEO.get(sp)):
            cand = [x for x in o_day if thu and x["sp"] == thu]
            if cand:
                hop = [x for x in cand if x["tang"] == tang]
                return [x["path"] for x in (hop or cand)]
    return []


def tim_kho():
    home = os.path.expanduser("~")
    for p in glob.glob(os.path.join(home, "OneDrive - *", "*" + KHO_TEN)):
        if os.path.isdir(p):
            return p
    return ""


def quet_kho(root):
    """Quet TEN file (khong mo file nao). -> dict gon de cache duoc."""
    idx, khong_doc, n_pdf, n_chuan, xmp = {}, [], 0, 0, []
    for r, _ds, fs in os.walk(root):
        for f in fs:
            if not f.lower().endswith(".pdf"):
                continue
            rel = os.path.relpath(os.path.join(r, f), root).replace(os.sep, "/")
            parts = rel.split("/")
            if f.upper().startswith("XMP"):
                xmp.append({"path": rel, "sp": sp_xmp(f), "tang": la_tang(rel)})
                continue
            n_pdf += 1
            yy_folder = parts[0][2:] if re.fullmatch(r"20\d{2}", parts[0]) else ""
            k = doc_ten(f[:-4], yy_folder)
            if not k:
                khong_doc.append(rel)
                continue
            yy, chu, so, e, chuan = k
            n_chuan += chuan
            idx.setdefault("%s|%s|%d|%d" % (yy, chu, so, e), []).append(
                {"path": rel, "sp": sp_cua_folder(parts[:-1]), "chuan": chuan})
    return {"idx": idx, "khong_doc": sorted(khong_doc), "n_pdf": n_pdf, "n_chuan": n_chuan,
            "xmp": xmp,
            "luc": dt.datetime.now().strftime("%H:%M:%S")}


def weburl(rel):
    return WEB.replace(" ", "%20") + "/" + "/".join(urllib.parse.quote(s) for s in rel.split("/"))


def tra_cuu(idx, pol, sp_ho_so):
    """-> (tinh_trang, [hit], ghi_chu). sp_ho_so = set ma san pham cua ho so (rong = khong biet)."""
    yy, chu, so, e = pol["yy"], pol["chu"], pol["so"], pol["e"]
    hits = idx.get("%s|%s|%d|%d" % (yy, chu, so, e), [])
    if not hits:
        co_e = sorted({int(k.split("|")[3]) for k in idx
                       if k.startswith("%s|%s|%d|" % (yy, chu, so))})
        if co_e:
            return ("Sai bản sửa đổi", [],
                    "Hồ sơ cần %s nhưng kho chỉ có %s, đề nghị P.TBHH cấp đúng bản"
                    % (ten_e(e), ", ".join(ten_e(x) for x in co_e)))
        return "Không có", [], "Kho không có file nào cho số đơn này"

    # Cung so, cung san pham ma co ca ban ten chuan lan ban ten cu (user copy ra ten moi, chua
    # xoa ban cu) -> lay ban ten chuan.
    theo_sp = {}
    for h in hits:
        theo_sp.setdefault(h["sp"], []).append(h)
    hits = [h for g in theo_sp.values()
            for h in ([x for x in g if x["chuan"]] or g)]

    note = []
    sps = {h["sp"] for h in hits}
    # San pham chi dung de CHON khi 1 so don co giay cua nhieu san pham. Chi co 1 san pham thi
    # khong co gi de chon, tra luon (ghi chu neu lech voi san pham cua ho so).
    if len(sps) > 1 and sp_ho_so:
        dung = [h for h in hits if h["sp"] in sp_ho_so]
        if dung:
            hits, sps = dung, {h["sp"] for h in dung}
    elif sp_ho_so and not (sps & sp_ho_so):
        note.append("GCN nằm ở folder %s, hồ sơ ghi %s"
                    % ("/".join(sorted(x or "không rõ" for x in sps)), "/".join(sorted(sp_ho_so))))
    if any(not h["chuan"] for h in hits):
        note.append("Tên file chưa theo chuẩn, tool đọc theo luật cũ")
    if len(hits) == 1:
        return "Tìm thấy", hits, "; ".join(note)
    if len(sps) > 1:
        note.insert(0, "Số đơn có GCN của %d sản phẩm, chọn đúng sản phẩm của hồ sơ" % len(sps))
        return "Nhiều sản phẩm", hits, "; ".join(note)
    note.insert(0, "Kho có %d bản cùng số, người kiểm tra tự chọn" % len(hits))
    return "Nhiều bản", hits, "; ".join(note)


def ten_e(e):
    return "E%02d" % e if e else "bản gốc"


def sp_tu_chuoi(s, bang):
    f = " " + re.sub(r"[^0-9A-Z]+", " ", fold(s)) + " "
    for code, keys in bang:
        if any(" %s " % k in f or (len(k) > 3 and k.replace(" ", "") in f.replace(" ", ""))
               for k in keys):
            return code
    return ""


def them_pol(ket, m, yc="", sp=""):
    k = m.group(0).upper()
    e = ket.setdefault(k, {"yy": m.group(1), "chu": m.group(4).upper(), "so": int(m.group(5)),
                           "e": int(m.group(6) or 0), "yc": [], "sp": set()})
    if yc and yc not in e["yc"]:
        e["yc"].append(yc)
    if sp:
        e["sp"].add(sp)


def dem_sheet(data):
    """-> [(ten sheet, so so don khac nhau)] de nguoi dung chon sheet lay so don."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ket = []
    for ws in wb.worksheets:
        pols = set()
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str):
                    pols.update(m.group(0).upper() for m in RE_POL.finditer(v))
        ket.append((ws.title, len(pols)))
    wb.close()
    return ket


def doc_file(data, sheets=None):
    """Quet MOI o cua cac sheet duoc chon (None = moi sheet): dong nao co so don thi lay, kem so
    yeu cau GSC va ten san pham cung dong. San pham con noi qua so GSC (sheet Data co Product
    nhung Policy trong, sheet Pias co Policy) nen bang GSC -> san pham van doc ca sheet KHONG chon."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    dong, gsc_sp = [], {}
    for ws in wb.worksheets:
        lay = sheets is None or ws.title in sheets
        for row in ws.iter_rows(values_only=True):
            pols, gscs, sp = [], [], ""
            for v in row:
                if v is None or isinstance(v, (int, float, dt.date)):
                    continue
                s = str(v)
                pols += list(RE_POL.finditer(s))
                gscs += RE_GSC.findall(s)
                if not sp and fold(s).lstrip().startswith(("BAO HANH", "BAO HIEM")):
                    sp = sp_tu_chuoi(s, SP_CLAIM)
            if sp:
                for g in gscs:
                    gsc_sp[g] = sp
            if pols and lay:
                dong.append((pols, gscs, sp))
    wb.close()
    ket = {}
    for pols, gscs, sp in dong:
        sps = {sp} if sp else {gsc_sp[g] for g in gscs if g in gsc_sp}
        for m in pols:
            if not gscs:
                them_pol(ket, m, "", next(iter(sps), "") if len(sps) == 1 else "")
            for g in gscs:
                them_pol(ket, m, g, gsc_sp.get(g, sp))
    return ket


def doc_text(txt):
    ket = {}
    for m in RE_POL.finditer(txt or ""):
        them_pol(ket, m)
    return ket


MAU = {"Tìm thấy": "E2EFDA", "Nhiều bản": "FFF2CC", "Nhiều sản phẩm": "FFF2CC",
       "Sai bản sửa đổi": "FCE4D6", "Không có": "EDEDED"}
THU_TU = ["Tìm thấy", "Nhiều sản phẩm", "Nhiều bản", "Sai bản sửa đổi", "Không có"]


def ke_hoach_xmp(rows):
    """XMP la bang ke ca thang, nhieu GCN dung chung 1 file -> trong ban in moi XMP chi xuat hien
    1 LAN, ngay sau GCN dau tien dung no (thu tu = thu tu dong ket qua).
    -> {so_don: (so_don_in_kem, so_don_dung_chung)} cho cac dong Tim thay co dung 1 XMP."""
    dau, dem, kh = {}, {}, {}
    for x in rows:
        if x["tt"] == "Tìm thấy" and len(x["xmp"]) == 1:
            r = x["xmp"][0]
            dau.setdefault(r, x["pol"])
            dem[r] = dem.get(r, 0) + 1
    for x in rows:
        if x["tt"] == "Tìm thấy" and len(x["xmp"]) == 1:
            r = x["xmp"][0]
            kh[x["pol"]] = (dau[r], dem[r])
    return kh


def tt_xmp(x):
    """Tinh trang XMP, cung bo chu + mau voi cot Tinh trang GCN. XMP chi do khi GCN ra dung 1 giay
    (ghep theo folder cua giay do), GCN chua chot thi chua biet XMP nao."""
    if x["tt"] != "Tìm thấy":
        return "Chưa xét"
    return "Tìm thấy" if len(x["xmp"]) == 1 else "Nhiều bản" if x["xmp"] else "Không có"


def xuat_excel(rows, n_link):
    hdr = (["STT", "Số hợp đồng", "Sản phẩm hồ sơ", "Số HS", "Số yêu cầu", "Tình trạng GCNBH"]
           + ["GCNBH %d" % (i + 1) for i in range(n_link)]
           + ["Tình trạng XMP", "XMP", "XMP trong bản in", "Ghi chú"])
    w = [5, 27, 10, 7, 36, 16] + [34] * n_link + [16, 34, 34, 60]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GCNBH"
    ws.sheet_view.showGridLines = False
    gach = Border(bottom=Side(style="thin", color="7F7F7F"))
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, name="Arial", size=10)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = gach
        ws.column_dimensions[get_column_letter(i)].width = w[i - 1]
    kh = ke_hoach_xmp(rows)
    for r, x in enumerate(rows, 2):
        vals = [r - 1, x["pol"], x["sp"], len(x["yc"]), ", ".join(x["yc"]), x["tt"]]
        vals += [""] * n_link + [tt_xmp(x), "", "", x["note"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(vertical="top", wrap_text=i in (5, len(vals)))
        ws.cell(row=r, column=6).fill = PatternFill("solid", start_color=MAU.get(x["tt"], "FFFFFF"))
        if tt_xmp(x) in MAU:
            ws.cell(row=r, column=7 + n_link).fill = PatternFill("solid", start_color=MAU[tt_xmp(x)])
        for j, h in enumerate(x["hits"][:n_link]):
            c = ws.cell(row=r, column=7 + j, value="%s · %s" % (h["sp"] or "?", h["path"].split("/")[-1]))
            c.hyperlink = weburl(h["path"])
            c.font = Font(name="Arial", size=10, color="0563C1", underline="single")
            c.alignment = Alignment(vertical="top")
        c = ws.cell(row=r, column=8 + n_link)
        if len(x["xmp"]) > 1:
            c.value = "%d XMP khả dĩ, tự chọn: %s" % (
                len(x["xmp"]), "; ".join(p.split("/")[-1] for p in x["xmp"]))
            c.hyperlink = weburl(x["xmp"][0])
            c.font = Font(name="Arial", size=10, color="0563C1", underline="single")
        elif x["xmp"]:
            c.value = x["xmp"][0].split("/")[-1]
            c.hyperlink = weburl(x["xmp"][0])
            c.font = Font(name="Arial", size=10, color="0563C1", underline="single")
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if x["pol"] in kh:
            dau, n = kh[x["pol"]]
            ws.cell(row=r, column=9 + n_link, value=(
                ("In kèm số đơn này, dùng chung %d số đơn" % n if n > 1 else "In kèm số đơn này")
                if dau == x["pol"] else "Đã in kèm %s" % dau))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(hdr)), len(rows) + 1)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ---- In tat ca: gop moi GCN "Tim thay" thanh 1 PDF de mo ra Ctrl+P mot lan -------------------
# File trong kho la placeholder OneDrive (0 byte). Doc file la OneDrive tu tai ve (~0,5 MB/giay),
# doc xong tra lai trang thai "chi online" de o C khong day dan (C chi 199 GB).
_PINNED, _UNPINNED = 0x80000, 0x100000


def _tra_lai_online(path):
    try:
        k32 = ctypes.windll.kernel32
        a = k32.GetFileAttributesW(path)
        if a != -1 and a != 0xFFFFFFFF:
            k32.SetFileAttributesW(path, (a & ~_PINNED) | _UNPINNED)
    except Exception:
        pass


def _doc_bytes(path):
    with open(path, "rb") as f:
        return f.read()


# File gop ~0,5 MB/trang (anh scan 200 dpi). Dot 61 = 92 trang / 45,7 MB: may in bo do nua sau
# (dung cho cac so dau 26 vi xep cuoi). Nen anh ve 150 dpi (45,7 -> 20,9 MB, chu + dau van ro)
# va tach file toi da N trang (user chon, mac dinh TRANG_MOI_FILE), cat dung ranh gioi giay de
# XMP di lien GCN dung no.
TRANG_MOI_FILE = 30
NEN_RONG, NEN_Q = 1240, 60


def _nen(data):
    """PDF bytes -> pymupdf doc, trang nao chi co 1 anh scan thi dung lai trang bang anh nho hon."""
    import pymupdf
    from PIL import Image
    src = pymupdf.open(stream=data, filetype="pdf")
    out = pymupdf.open()
    for pg in src:
        ims = pg.get_images(full=True)
        try:
            if len(ims) != 1:
                raise ValueError
            img = Image.open(io.BytesIO(src.extract_image(ims[0][0])["image"])).convert("RGB")
            if img.width > NEN_RONG:
                img = img.resize((NEN_RONG, round(img.height * NEN_RONG / img.width)), Image.LANCZOS)
            b = io.BytesIO()
            img.save(b, "JPEG", quality=NEN_Q, optimize=True)
            moi = out.new_page(width=pg.rect.width, height=pg.rect.height)
            moi.insert_image(moi.rect, stream=b.getvalue())
        except Exception:
            out.insert_pdf(src, from_page=pg.number, to_page=pg.number)
    src.close()
    return out


def gop_pdf(root, rows, kem_xmp=True, progress=None, trang_moi_file=TRANG_MOI_FILE):
    """-> ([(pdf_bytes, so_giay, so_trang)], so_giay, so_trang, [loi]). Chi lay dong 'Tim thay'
    (dung 1 giay). XMP chi in khi ghep duoc DUNG 1 file; 2 ung vien tro len thi bo qua va bao."""
    import pymupdf
    viec = [x for x in rows if x["tt"] == "Tìm thấy"]
    can = []
    for x in viec:
        can.append(x["hits"][0]["path"])
        if kem_xmp and len(x["xmp"]) == 1:
            can.append(x["xmp"][0])
    can = list(dict.fromkeys(can))
    full = {r: os.path.join(root, *r.split("/")) for r in can}
    loi, data = [], {}
    with ThreadPoolExecutor(max_workers=6) as ex:
        futs = {r: ex.submit(_doc_bytes, full[r]) for r in can}
        for i, r in enumerate(can, 1):
            if progress:
                progress(i - 1, len(can), r.split("/")[-1])
            try:
                data[r] = futs[r].result(timeout=180)
            except Exception as exc:
                loi.append("Không tải được %s (%s)" % (r, exc))
            finally:
                _tra_lai_online(full[r])

    if progress:
        progress(len(can), len(can), "nén ảnh…")
    docs = {}
    for r in data:
        try:
            docs[r] = _nen(data[r])
        except Exception as exc:
            loi.append("Không đọc được %s (%s)" % (r, exc))

    # Moi "giay" = GCN + XMP in kem (neu co) -> khong bao gio bi cat doi giua 2 file
    giay, da_in = [], set()
    for x in viec:
        h = x["hits"][0]
        if h["path"] not in docs:
            continue
        g = [(h["path"], "%s · %s" % (x["pol"], h["sp"] or "?"), 1)]
        xr = x["xmp"][0] if kem_xmp and len(x["xmp"]) == 1 else None
        if xr and xr not in da_in and xr in docs:
            da_in.add(xr)
            g.append((xr, "XMP · " + xr.split("/")[-1][:-4], 2))
        giay.append(g)

    phan, cur, n_trang = [], [], 0
    for g in giay:
        t = sum(docs[r].page_count for r, _, _ in g)
        if cur and n_trang + t > trang_moi_file:
            phan.append(cur)
            cur, n_trang = [], 0
        cur.append(g)
        n_trang += t
    if cur:
        phan.append(cur)

    ket = []
    for nhom in phan:
        out, toc = pymupdf.open(), []
        for g in nhom:
            for r, nhan, cap in g:
                toc.append([cap, nhan, out.page_count + 1])
                out.insert_pdf(docs[r])
        out.set_toc(toc)
        ket.append((out.tobytes(garbage=3, deflate=True), len(nhom), out.page_count))
        out.close()
    if progress:
        progress(len(can), len(can), "xong")
    return ket, len(giay), sum(k[2] for k in ket), loi


def chay(ket, idx, chon, goi_y_file, xmps=()):
    rows = []
    for pol, e in sorted(ket.items()):
        if chon == TU_NHAN:
            sp = e["sp"] or ({goi_y_file} if goi_y_file else set())
        elif chon == TAT_CA:
            sp = set()
        else:
            sp = {chon}
        tt, hits, note = tra_cuu(idx, e, sp)
        rows.append({"pol": pol, "sp": "/".join(sorted(sp)), "yc": e["yc"], "tt": tt,
                     "hits": hits, "note": note,
                     "xmp": tim_xmp(xmps, hits[0]) if len(hits) == 1 else []})
    return rows
